"""Importação de extratos de corretoras.

Cada linha gera uma impressão digital (`external_id`). Carregar o mesmo extrato
várias vezes, ou extratos com períodos sobrepostos, só importa as linhas novas.
"""
import csv
import hashlib
import io
import re
from datetime import datetime
from pathlib import Path

from ..models import Broker, FxRate, ImportBatch, Instrument, Price, Transaction, db

TYPES = ("BUY", "SELL", "DIVIDEND", "DEPOSIT", "WITHDRAWAL", "FEE", "INTEREST", "TRANSFER_IN", "TRANSFER_OUT",
         "CONVERT_IN", "CONVERT_OUT")

TYPE_ALIASES = {
    "buy": "BUY", "compra": "BUY", "sell": "SELL", "venda": "SELL",
    "dividend": "DIVIDEND", "dividendo": "DIVIDEND", "deposit": "DEPOSIT", "depósito": "DEPOSIT",
    "deposito": "DEPOSIT", "withdrawal": "WITHDRAWAL", "levantamento": "WITHDRAWAL",
    "fee": "FEE", "comissão": "FEE", "comissao": "FEE", "interest": "INTEREST", "juros": "INTEREST",
    "transfer_in": "TRANSFER_IN", "transfer_out": "TRANSFER_OUT", "transferência entrada": "TRANSFER_IN",
    "transferência saída": "TRANSFER_OUT",
}

# Corretoras apresentadas na página de importação. `parser` None = ainda não suportado.
BROKERS = [
    {"key": "generic", "name": "CSV genérico", "format": "CSV", "parser": "generic", "color": "#64748b"},
    {"key": "trading212", "name": "Trading 212", "format": "Extrato CSV", "parser": "trading212", "color": "#0ea5e9"},
    {"key": "degiro", "name": "DEGIRO", "format": "Extrato CSV", "parser": None, "color": "#38bdf8"},
    {"key": "xtb", "name": "XTB", "format": "Extrato XLSX", "parser": "xtb", "color": "#ef4444"},
    {"key": "revolut", "name": "Revolut", "format": "Extrato PDF", "parser": None, "color": "#e2e8f0"},
    {"key": "traderepublic", "name": "Trade Republic", "format": "Extrato CSV", "parser": "traderepublic",
     "color": "#f8fafc"},
    {"key": "ibkr", "name": "IBKR", "format": "Extrato CSV", "parser": None, "color": "#dc2626"},
    {"key": "etoro", "name": "eToro", "format": "Extrato PDF", "parser": None, "color": "#22c55e"},
]

# logótipos em static/icons/brokers/<chave>.svg|png; a chave é deduzida do nome da corretora,
# para servir também às corretoras criadas na importação com um nome próprio (ex.: "XTB – conta EUR")
_LOGO_DIR = Path(__file__).resolve().parent.parent / "static" / "icons" / "brokers"
_LOGO_NAMES = (("trading212", ("trading 212", "trading212")), ("xtb", ("xtb",)),
               ("traderepublic", ("trade republic", "traderepublic")))


def broker_logo(name_or_key):
    """Caminho (relativo a /static) do logótipo da corretora, ou None se não houver."""
    text = (name_or_key or "").strip().lower()
    for key, names in _LOGO_NAMES:
        if text == key or any(n in text for n in names):
            for ext in ("svg", "png"):
                if (_LOGO_DIR / f"{key}.{ext}").is_file():
                    return f"icons/brokers/{key}.{ext}"
    return None


INSTRUCTIONS = {
    "generic": [
        "Cria um ficheiro CSV com o cabeçalho: <code>date,type,ticker,quantity,price,amount,fees,tax,currency,name</code>.",
        "<code>type</code>: BUY, SELL, DIVIDEND, DEPOSIT, WITHDRAWAL, FEE ou INTEREST (também em português).",
        "Datas em <code>AAAA-MM-DD</code> ou <code>DD/MM/AAAA</code>; decimais com ponto ou vírgula.",
        "Para compras/vendas preenche <code>quantity</code> e <code>price</code>; para movimentos de caixa usa <code>amount</code>.",
    ],
    "xtb": [
        "Inicia sessão no <b>xStation</b> da XTB (web ou aplicação de computador).",
        "Abre <b>Histórico da conta</b> e escolhe o período <b>desde a abertura da conta</b>.",
        "Clica em <b>Exportar</b> e escolhe o formato <b>XLSX</b>: o ficheiro traz as folhas <i>Cash Operations</i>, <i>Closed Positions</i> e <i>Open Positions</i>.",
        "Carrega aqui o ficheiro. Os planos de investimento (<i>Investment Plans</i>) e as transferências entre subcontas são tratados automaticamente.",
    ],
    "traderepublic": [
        "Abre a aplicação ou o site da <b>Trade Republic</b> e entra no <b>Perfil</b>.",
        "Em <b>Atividade</b> / <b>Documentos</b>, escolhe <b>Exportar transações</b> (formato <b>CSV</b>).",
        "Escolhe o período desde a abertura da conta: o ficheiro chama-se <code>transactions_AAAA-MM-DD_AAAA-MM-DD.csv</code>.",
        "Carrega aqui o ficheiro. Compras, vendas, planos de poupança, juros, dividendos e transferências são importados.",
    ],
    "trading212": [
        "Inicia sessão na tua conta <b>Trading 212</b> em app.trading212.com.",
        "Clica no teu <b>nome</b> (canto superior direito) e abre o <b>Histórico</b>.",
        "Clica em <b>Exportar</b> e liga todos os tipos de dados (<b>Ordens</b>, <b>Dividendos</b>, <b>Juros</b> e <b>Transações</b>).",
        "Escolhe o período desde a tua <b>primeira transação</b>. A Trading 212 só exporta 1 ano de cada vez; repete para cada ano e carrega todos os ficheiros.",
    ],
}


def _num(v):
    if v is None:
        return 0.0
    v = str(v).strip().replace(" ", "").replace(" ", "")
    if not v or v.lower() in {"nan", "-"}:
        return 0.0
    if "," in v and "." in v:
        v = v.replace(".", "").replace(",", ".") if v.rfind(",") > v.rfind(".") else v.replace(",", "")
    elif "," in v:
        v = v.replace(",", ".")
    return float(v)


def _date(v):
    v = v.strip()[:19]
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"data inválida: {v!r}")


def _fingerprint(*parts):
    return hashlib.sha1("|".join(str(p) for p in parts).encode()).hexdigest()


class ImportError_(ValueError):
    """Erro de importação com mensagem para o utilizador."""


def _guess_kind(name):
    n = (name or "").upper()
    return "etf" if any(k in n for k in ("(ACC)", "(DIST)", " ETF", "UCITS")) else "stock"


_ISIN = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")


def _rename_instrument(old, new):
    """Passa todas as referências de um ativo para outro ticker (ex.: ISIN -> ticker real)."""
    from ..models import Alert, Event, News, WatchItem
    for model, col in ((Transaction, "instrument"), (Price, "instrument"), (Event, "instrument"),
                       (News, "instrument"), (Alert, "instrument"), (WatchItem, "instrument")):
        model.query.filter(getattr(model, col) == old.ticker).update({col: new.ticker}, synchronize_session=False)
    db.session.delete(old)
    db.session.flush()


def _instrument(ticker, name=None, currency="EUR", isin=None, symbol=None, kind=None):
    """Devolve o ticker do ativo, criando-o se necessário.

    O mesmo título comprado em corretoras diferentes é unificado pelo ISIN quando uma delas só
    identifica o ativo pelo ISIN (ex.: Trade Republic): partilham posição e cotações.
    """
    ticker = ticker.strip().upper()
    inst = db.session.get(Instrument, ticker)
    if inst is None and isin:
        same = Instrument.query.filter_by(isin=isin).all()
        if _ISIN.match(ticker) and same:  # ex.: Trade Republic usa o ISIN como identificador
            return next((i for i in same if i.currency == currency), same[0]).ticker
        placeholder = next((i for i in same if _ISIN.match(i.ticker)), None)
        if placeholder:  # já existia com o ISIN como ticker: passa a usar o ticker real
            new = Instrument(ticker=ticker, instrument_name=name or placeholder.instrument_name, currency=currency,
                             isin=isin, kind=kind or placeholder.kind, color=placeholder.color,
                             data_symbol=symbol or yahoo_symbol(ticker, currency) or placeholder.data_symbol)
            db.session.add(new)
            db.session.flush()
            _rename_instrument(placeholder, new)
            return ticker
    if inst is None:
        db.session.add(Instrument(ticker=ticker, instrument_name=name or ticker, currency=currency, isin=isin,
                                  kind=kind or _guess_kind(name), data_symbol=symbol or yahoo_symbol(ticker, currency)))
        db.session.flush()
    elif isin and not inst.isin:
        inst.isin = isin
    return ticker


def _read(text):
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def parse_generic(rows):
    for r in rows:
        r = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
        typ = TYPE_ALIASES.get(r.get("type", "").lower(), r.get("type", "").upper())
        if typ not in TYPES:
            raise ValueError(f"tipo desconhecido: {r.get('type')!r}")
        yield {
            "date": _date(r["date"]), "type": typ, "ticker": r.get("ticker") or None,
            "name": r.get("name"), "quantity": _num(r.get("quantity")), "price": _num(r.get("price")),
            "amount": _num(r.get("amount")), "fees": _num(r.get("fees")), "tax": _num(r.get("tax")),
            "currency": (r.get("currency") or "EUR").upper(), "external_id": r.get("id") or None,
            "result": _num(r["result"]) if r.get("result") else None, "note": r.get("note") or None,
            "isin": r.get("isin") or None,
        }


T212_FEES = ("Currency conversion fee", "French transaction tax", "Stamp duty", "Stamp duty reserve tax",
             "Finra fee", "Transaction fee")


def yahoo_symbol(ticker, currency):
    """Símbolo provável no Yahoo Finance (as ações europeias resolvem-se pelo ISIN ao atualizar)."""
    t = ticker.strip().upper().replace("/", "-")
    return {"GBP": t + ".L", "CAD": t + ".TO"}.get(currency, t if currency == "USD" else None)


def parse_trading212(rows):
    """Extrato CSV da Trading 212 (conta em EUR, com ou sem subconta em USD).

    Notas sobre o formato:
      * "Exchange rate" nas compras/vendas = moeda do ativo por unidade da moeda do "Total";
      * nos dividendos, "Price / share" é o dividendo líquido por ação e "Total" o valor líquido;
      * preços em GBX (pence) são convertidos para GBP.
    """
    rows = list(rows)
    # conversões de fundos: "Stock distribution" do novo ativo + "venda" a preço 0 do antigo no mesmo dia
    distribution_days = {_date(r.get("Time") or r.get("Time (UTC)") or "") for r in rows
                         if (r.get("Action") or "").strip().lower() == "stock distribution"}
    for r in rows:
        action = (r.get("Action") or "").strip().lower()
        d = _date(r.get("Time") or r.get("Time (UTC)") or "")
        total = _num(r.get("Total"))
        total_cur = (r.get("Currency (Total)") or "EUR").upper()
        price_cur = (r.get("Currency (Price / share)") or total_cur).upper()
        factor = 1.0
        if price_cur == "GBX":  # pence -> libras
            price_cur, factor = "GBP", 0.01
        qty = _num(r.get("No. of shares"))
        price = _num(r.get("Price / share")) * factor
        rate = (_num(r.get("Exchange rate")) or 1.0) * factor
        ticker = (r.get("Ticker") or "").strip() or None
        base = {"date": d, "external_id": r.get("ID") or None, "ticker": ticker, "name": r.get("Name"),
                "isin": r.get("ISIN") or None, "quantity": 0.0, "price": 0.0, "amount": 0.0, "fees": 0.0,
                "tax": 0.0, "currency": total_cur, "result": None, "note": None}

        def to_price_cur(value, cur):
            cur = (cur or total_cur).upper()
            if cur == "GBX":
                value, cur = value / 100, "GBP"
            return value if cur == price_cur else value * rate

        if action.endswith(" sell") and d in distribution_days and _num(r.get("Price / share")) == 0 \
                and _num(r.get("Total")) == 0:
            result = _num(r.get("Result")) if r.get("Result") else None
            yield {**base, "type": "CONVERT_OUT", "quantity": qty, "currency": price_cur,
                   "result": to_price_cur(result, r.get("Currency (Result)")) if result is not None else None,
                   "note": "Conversão de fundo"}
        elif action == "stock distribution":
            yield {**base, "type": "CONVERT_IN", "quantity": qty, "currency": price_cur, "note": "Conversão de fundo"}
        elif action == "stock split close":
            yield {**base, "type": "CONVERT_OUT", "quantity": qty, "currency": price_cur, "note": "Split"}
        elif action == "stock split open":
            yield {**base, "type": "CONVERT_IN", "quantity": qty, "currency": price_cur, "note": "Split",
                   "_price": price}
        elif action == "stock dividends":
            yield {**base, "type": "BUY", "quantity": qty, "price": 0.0, "currency": price_cur,
                   "note": "Dividendo em ações"}
        elif action == "dividend adjustment":  # acerto da retenção na fonte, sem ativo associado
            yield {**base, "ticker": None, "type": "INTEREST", "amount": total,
                   "note": r.get("Notes") or "Acerto de dividendos"}
        elif action.endswith(" buy") or action.endswith(" sell"):
            fees = sum(to_price_cur(abs(_num(r.get(c))), r.get(f"Currency ({c})")) for c in T212_FEES)
            result = None
            if action.endswith("sell") and r.get("Result"):
                result = to_price_cur(_num(r.get("Result")), r.get("Currency (Result)"))
            obs = {"_price": price}
            # linhas residuais (total 0) trazem taxa 1.0 por omissão: não servem como câmbio
            if total_cur == "EUR" and price_cur != "EUR" and total and abs(rate - factor) > 1e-9:
                obs["_fx"] = (price_cur, rate)
            yield {**base, **obs, "type": "BUY" if action.endswith("buy") else "SELL", "quantity": qty,
                   "price": price, "currency": price_cur, "fees": fees, "result": result}
        elif action.startswith("dividend"):
            tax = abs(_num(r.get("Withholding tax")))
            tax_cur = (r.get("Currency (Withholding tax)") or price_cur).upper()
            net_local = qty * price
            if net_local > 0 and total > 0:
                eff = total / net_local  # moeda do Total por unidade da moeda do ativo
                if tax_cur != price_cur:
                    tax = tax / eff
                obs = {"_fx": (price_cur, 1 / eff)} if total_cur == "EUR" and price_cur != "EUR" else {}
                yield {**base, **obs, "type": "DIVIDEND", "quantity": qty, "amount": net_local + tax,
                       "tax": tax, "currency": price_cur}
            else:
                yield {**base, "type": "DIVIDEND", "quantity": qty, "amount": total + tax, "tax": tax}
        elif action == "deposit":
            yield {**base, "type": "DEPOSIT", "amount": total}
        elif action == "withdrawal":
            yield {**base, "type": "WITHDRAWAL", "amount": abs(total)}
        elif action == "card debit":
            merchant = r.get("Merchant name")
            yield {**base, "type": "WITHDRAWAL", "amount": abs(total),
                   "note": f"Cartão: {merchant}" if merchant else "Cartão"}
        elif "interest" in action:
            yield {**base, "type": "INTEREST", "amount": total, "note": r.get("Notes") or None}
        elif action in ("transfer out", "transfer in"):
            yield {**base, "type": "TRANSFER_OUT" if action == "transfer out" else "TRANSFER_IN",
                   "quantity": qty, "price": price, "currency": price_cur, "_price": price}
        elif action == "spin off":
            # ações recebidas sem custo
            yield {**base, "type": "BUY", "quantity": qty, "price": 0.0, "currency": price_cur, "note": "Spin-off"}
        elif action == "currency conversion":
            fee = abs(_num(r.get("Currency conversion fee")))
            if fee:
                yield {**base, "ticker": None, "type": "FEE", "amount": fee,
                       "currency": (r.get("Currency (Currency conversion fee)") or total_cur).upper(),
                       "note": r.get("Notes") or "Conversão de moeda"}
            else:
                yield {**base, "type": None, "_ignored": r.get("Action")}
        elif action.endswith("fee"):  # ADR Fee e outras comissões
            yield {**base, "ticker": None, "type": "FEE", "amount": abs(total),
                   "note": f"{r.get('Action')} {ticker or ''}".strip()}
        else:
            yield {**base, "type": None, "_ignored": r.get("Action")}


def parse_traderepublic(rows):
    """Exportação CSV de transações da Trade Republic.

    Os ativos são identificados pelo ISIN (coluna "symbol"); "amount" é o valor bruto do movimento e
    "fee"/"tax" vêm separados (o imposto nas compras/vendas é imposto de transação, ex.: francês).
    """
    for r in rows:
        typ = (r.get("type") or "").strip().upper()
        category = (r.get("category") or "").strip().upper()
        d = _date((r.get("date") or r.get("datetime") or "")[:10])
        amount = _num(r.get("amount"))
        fee, tax = abs(_num(r.get("fee"))), abs(_num(r.get("tax")))
        isin = (r.get("symbol") or "").strip().upper() or None
        cur = (r.get("currency") or "EUR").upper()
        base = {"date": d, "external_id": f"tr-{r.get('transaction_id')}" if r.get("transaction_id") else None,
                "ticker": isin, "name": r.get("name") or None, "isin": isin, "quantity": 0.0, "price": 0.0,
                "amount": 0.0, "fees": 0.0, "tax": 0.0, "currency": cur, "result": None,
                "note": r.get("description") or None,
                "_kind": "etf" if (r.get("asset_class") or "").upper() in ("FUND", "ETF") else "stock"}
        if typ in ("BUY", "SELL"):
            qty, price = abs(_num(r.get("shares"))), _num(r.get("price"))
            if not price and qty:
                price = abs(amount) / qty
            yield {**base, "type": typ, "quantity": qty, "price": price, "fees": fee + tax, "_price": price}
        elif "DIVIDEND" in typ or typ in ("DISTRIBUTION", "CAPITAL_RETURN"):
            yield {**base, "type": "DIVIDEND", "quantity": abs(_num(r.get("shares"))), "amount": amount + tax,
                   "tax": tax}
        elif "INTEREST" in typ:
            yield {**base, "ticker": None, "type": "INTEREST", "amount": amount + tax, "tax": tax}
        elif category == "CASH" or "TRANSFER" in typ or "CARD" in typ or "PAYMENT" in typ:
            if not amount:
                yield {**base, "type": None, "_ignored": typ or "(vazio)"}
            else:
                yield {**base, "ticker": None, "type": "DEPOSIT" if amount > 0 else "WITHDRAWAL",
                       "amount": abs(amount)}
        elif "FEE" in typ:
            yield {**base, "ticker": None, "type": "FEE", "amount": abs(amount) or fee}
        else:
            yield {**base, "type": None, "_ignored": typ or "(vazio)"}


XTB_SUFFIX = {".FR": ".PA", ".NL": ".AS", ".PT": ".LS", ".ES": ".MC", ".UK": ".L", ".IT": ".MI", ".BE": ".BR",
              ".US": "", ".DE": ".DE", ".PL": ".WA", ".CH": ".SW", ".DK": ".CO", ".SE": ".ST", ".NO": ".OL",
              ".FI": ".HE", ".AT": ".VI", ".IE": ".IR"}
_XTB_TRADE = re.compile(r"(OPEN|CLOSE)\s+BUY\s+(?:\S+\s+)?([\d.]+)(?:/[\d.]+)?\s*@\s*([\d.]+)")


def xtb_yahoo_symbol(ticker):
    base, dot, suffix = ticker.rpartition(".")
    if not dot:
        return ticker
    return base + XTB_SUFFIX.get("." + suffix.upper(), "." + suffix.upper())


def _xtb_sheet(wb, name, first_header):
    """Linhas (dicts) de uma folha da XTB a partir da linha de cabeçalho que começa por `first_header`."""
    if name not in wb.sheetnames:
        return []
    rows = [r for r in wb[name].iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    idx = next((i for i, r in enumerate(rows) if r and r[0] == first_header), None)
    if idx is None:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[idx]]
    out = []
    for r in rows[idx + 1:]:
        if r[0] in ("Total", None) or (isinstance(r[0], str) and r[0] == first_header):
            break
        out.append(dict(zip(header, r)))
    return out


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    return _date(str(v))


def parse_xtb(data):
    """Extrato XLSX da XTB (conta em EUR).

    A folha "Cash Operations" tem todos os movimentos de dinheiro; quantidade e preço das
    compras/vendas vêm do comentário ("OPEN BUY VWCE.DE 5/5 @ 120.00"). A folha
    "Open Positions" dá o preço atual de cada posição na data do relatório.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("falta o pacote openpyxl (pip install openpyxl)") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ValueError("não é um ficheiro XLSX válido da XTB") from exc
    ops = _xtb_sheet(wb, "Cash Operations", "Type")
    if not ops:
        raise ValueError('não encontrei a folha "Cash Operations"')
    names = {}
    for op in ops:
        typ = str(op.get("Type") or "").strip().lower()
        amount = float(op.get("Amount") or 0)
        d = _as_date(op.get("Time"))
        ticker = (op.get("Ticker") or "").strip() or None
        comment = str(op.get("Comment") or "")
        base = {"date": d, "external_id": f"xtb-{op.get('ID')}" if op.get("ID") else None, "ticker": ticker,
                "name": op.get("Instrument") or None, "isin": None, "quantity": 0.0, "price": 0.0, "amount": 0.0,
                "fees": 0.0, "tax": 0.0, "currency": "EUR", "result": None, "note": None}
        if ticker:
            names[ticker] = (op.get("Instrument"), op.get("Category"))
            base["_symbol"] = xtb_yahoo_symbol(ticker)
            base["_kind"] = "etf" if op.get("Category") == "ETF" else "stock"
        if typ in ("stock purchase", "stock sell", "stock sale"):
            m = _XTB_TRADE.search(comment)
            if not m:
                yield {**base, "type": None, "_ignored": f"{op.get('Type')} sem quantidade/preço"}
                continue
            qty, price = float(m.group(2)), float(m.group(3))
            gross = qty * price
            # instrumentos noutra moeda: a XTB converte para EUR; usa o preço efetivo em EUR
            if gross and abs(abs(amount) - gross) / gross > 0.05:
                price, gross = abs(amount) / qty, abs(amount)
            buy = typ == "stock purchase"
            fees = round(abs(amount) - gross, 2) if buy else round(gross - abs(amount), 2)
            if fees <= 0:  # arredondamento ao cêntimo: o preço efetivo reproduz o valor exato
                price, fees = abs(amount) / qty, 0.0
            yield {**base, "type": "BUY" if buy else "SELL", "quantity": qty, "price": price,
                   "fees": fees, "_price": float(m.group(3)),
                   "note": op.get("Product") if op.get("Product") == "Investment Plans" else None}
        elif typ == "deposit":
            yield {**base, "ticker": None, "type": "DEPOSIT", "amount": abs(amount)}
        elif typ == "withdrawal":
            yield {**base, "ticker": None, "type": "WITHDRAWAL", "amount": abs(amount)}
        elif typ == "transfer":  # de/para outra conta XTB
            yield {**base, "ticker": None, "type": "DEPOSIT" if amount > 0 else "WITHDRAWAL", "amount": abs(amount),
                   "note": comment}
        elif typ == "subaccount transfer":
            yield {**base, "type": None, "_ignored": "Transferência entre subcontas"}
        elif typ == "free funds interest":
            yield {**base, "ticker": None, "type": "INTEREST", "amount": amount, "note": comment}
        elif typ == "free funds interest tax":
            yield {**base, "ticker": None, "type": "INTEREST", "amount": 0.0, "tax": abs(amount), "note": comment}
        elif typ.startswith("divid"):
            yield {**base, "type": "DIVIDEND", "amount": amount}
        elif "withholding" in typ:
            yield {**base, "type": "DIVIDEND", "amount": 0.0, "tax": abs(amount), "note": comment}
        elif "fee" in typ or "commission" in typ:
            yield {**base, "ticker": None, "type": "FEE", "amount": abs(amount), "note": comment}
        elif typ == "close trade":  # resultado de um CFD (não é uma posição em carteira)
            yield {**base, "ticker": None, "type": "INTEREST" if amount >= 0 else "FEE", "amount": abs(amount),
                   "note": f"Resultado CFD {ticker or ''}: {comment}".strip()}
        else:
            yield {**base, "type": None, "_ignored": op.get("Type")}

    yield from _xtb_transfers_in(wb, ops, _xtb_period_start(wb))

    # preço atual das posições abertas, na data do relatório
    if "Open Positions" in wb.sheetnames:
        report = next((r for r in wb["Open Positions"].iter_rows(values_only=True)
                       if r and str(r[0]).startswith("Data as of")), None)
        report_date = _as_date(report[1]) if report and report[1] else None
        for r in _xtb_positions(wb) if report_date else []:
            if r.get("Ticker") and r.get("Current price"):
                name, cat = names.get(r["Ticker"], (None, None))
                yield {"_quote": {"ticker": r["Ticker"], "name": name, "date": report_date,
                                  "price": float(r["Current price"]), "symbol": xtb_yahoo_symbol(r["Ticker"]),
                                  "kind": "etf" if cat == "ETF" else "stock"}}


def _pid(v):
    """Identificador de posição normalizado (a XTB usa texto numa folha e número noutra)."""
    if v in (None, ""):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _xtb_period_start(wb):
    """Início do período do extrato ("Date from (UTC)" na folha Cash Operations), se existir."""
    for r in wb["Cash Operations"].iter_rows(max_row=10, values_only=True):
        if r and r[0] == "Date from (UTC)" and isinstance(r[1], datetime):
            return r[1]
    return None


def _xtb_transfers_in(wb, ops, start=None):
    """Ações recebidas de outra corretora: posições com venda/abertas mas sem compra em dinheiro.

    Num extrato de um período curto, as compras anteriores ao início do período não aparecem: essas
    posições ficam marcadas (`_before_period`) e `import_file` ignora-as se já houver histórico importado.
    """
    purchased = {_pid(o.get("Position ID")) for o in ops if str(o.get("Type", "")).lower() == "stock purchase"}
    sells = {}
    for o in ops:
        if str(o.get("Type", "")).lower() in ("stock sell", "stock sale") and isinstance(o.get("Time"), datetime):
            sells.setdefault(_pid(o.get("Position ID")), []).append(o["Time"])
    rows = []
    for r in _xtb_sheet(wb, "Closed Positions", "Instrument"):
        pid, closed_at = _pid(r.get("Position ID")), r.get("Close Time (UTC)")
        # só conta a parte fechada com dinheiro (fechos técnicos, sem venda, não são transferências)
        if pid not in purchased and isinstance(closed_at, datetime) and any(
                abs((t - closed_at).total_seconds()) < 5 for t in sells.get(pid, [])):
            rows.append((r["Ticker"], r.get("Instrument"), r.get("Category"), float(r["Volume"]),
                         float(r["Open Price"]), r["Open Time (UTC)"], pid))
    if "Open Positions" in wb.sheetnames:
        for r in _xtb_positions(wb):
            pid = _pid(r.get("Instrument/Position"))
            if r.get("Type") == "BUY" and pid and pid not in purchased and r.get("Open price"):
                rows.append((r["Ticker"], None, None, float(r["Volume"]), float(r["Open price"]),
                             r.get("Open time (UTC)"), pid))
    for ticker, name, cat, vol, price, opened, pid in rows:
        d = _as_date(opened)
        before = bool(start and isinstance(opened, datetime) and opened < start)
        yield {"date": d, "external_id": f"xtb-in-{pid}-{vol}", "ticker": ticker, "name": name, "isin": None,
               "quantity": vol, "price": price, "amount": 0.0, "fees": 0.0, "tax": 0.0, "currency": "EUR",
               "result": None, "note": "Ações recebidas de outra corretora", "type": "TRANSFER_IN",
               "_symbol": xtb_yahoo_symbol(ticker), "_kind": "etf" if cat == "ETF" else None, "_price": price,
               "_before_period": before}


def _xtb_positions(wb):
    """Linhas de posições abertas (a folha tem um resumo antes da tabela de posições)."""
    rows = [r for r in wb["Open Positions"].iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    idx = next((i for i, r in enumerate(rows) if r and "Ticker" in r and "Current price" in r), None)
    if idx is None:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[idx]]
    return [dict(zip(header, r)) for r in rows[idx + 1:] if r[0] not in (None, "Total")]


PARSERS = {"generic": parse_generic, "trading212": parse_trading212, "traderepublic": parse_traderepublic}
BINARY_PARSERS = {"xtb": parse_xtb}


def detect_broker(filename, data):
    """Corretora provável de um extrato, pelo conteúdo (para importar sem escolher a corretora)."""
    if data[:2] == b"PK":  # XLSX
        return "xtb"
    lines = data[:4096].decode("utf-8-sig", errors="replace").splitlines() if data else []
    cols = {c.strip().strip('"').lower() for c in re.split(r"[,;\t]", lines[0])} if lines else set()
    if {"transaction_id", "asset_class"} <= cols:
        return "traderepublic"
    if "action" in cols and ({"time", "time (utc)"} & cols):
        return "trading212"
    return "generic"


def get_or_create_broker(name, color="#3b82f6"):
    b = Broker.query.filter_by(broker_name=name).first()
    if b is None:
        b = Broker(broker_name=name, color=color)
        db.session.add(b)
        db.session.flush()
    return b.id


def import_file(broker_key, filename, data, broker_name=None):
    """Importa um ficheiro e devolve o `ImportBatch` criado."""
    info = next((b for b in BROKERS if b["key"] == broker_key), None)
    if not info or not info["parser"]:
        raise ImportError_("Esta corretora ainda não é suportada.")
    try:
        if info["parser"] in BINARY_PARSERS:
            items = list(BINARY_PARSERS[info["parser"]](data))
        else:
            if data[:2] == b"PK":
                raise ValueError("este ficheiro é um XLSX; escolhe a corretora certa ou exporta em CSV")
            items = list(PARSERS[info["parser"]](_read(data.decode("utf-8-sig", errors="replace"))))
    except (KeyError, ValueError) as exc:
        raise ImportError_(f"Não foi possível ler o ficheiro: {exc}") from exc

    broker = get_or_create_broker(broker_name or info["name"], info["color"])
    batch = ImportBatch(broker=broker, source=broker_key, filename=filename, imported=0, skipped=0)
    db.session.add(batch)
    db.session.flush()

    existing = {e for (e,) in db.session.query(Transaction.external_id).filter_by(broker=broker)}
    has_history = bool(existing)
    dates, fx_obs, price_obs = [], {}, {}
    ignored = {}
    for n, item in enumerate(items, start=2):
        if item.get("_quote"):  # cotação observada no extrato (ex.: preço atual da posição)
            q = item["_quote"]
            inst = _instrument(q["ticker"], q.get("name"), q.get("currency", "EUR"), symbol=q.get("symbol"),
                               kind=q.get("kind"))
            price_obs[(inst, q["date"])] = (q["price"], q.get("currency", "EUR"))
            continue
        if item.pop("_before_period", False) and has_history:
            # comprada antes do período do extrato: a compra já veio numa importação anterior
            item = {"_ignored": "Posições compradas antes do período do extrato (já importadas)"}
        if item.get("type") is None:
            ignored[item.get("_ignored")] = ignored.get(item.get("_ignored"), 0) + 1
            continue
        fx = item.pop("_fx", None)
        item.pop("_price", None)  # os preços das transações são usados ao carregar a carteira
        symbol, kind = item.pop("_symbol", None), item.pop("_kind", None)
        fp = item.pop("external_id") or _fingerprint(
            item["date"], item["type"], item["ticker"], round(item["quantity"], 6),
            round(item["price"], 6), round(item["amount"], 2))
        dates.append(item["date"])
        ticker, name_, isin = item.pop("ticker"), item.pop("name"), item.pop("isin", None)
        if fx:
            fx_obs.setdefault((fx[0], item["date"]), []).append(fx[1])
        if fp in existing:
            batch.skipped += 1
            continue
        existing.add(fp)
        instrument = None
        if item["type"] in ("BUY", "SELL", "DIVIDEND", "TRANSFER_IN", "TRANSFER_OUT", "CONVERT_IN", "CONVERT_OUT"):
            if not ticker:
                db.session.rollback()
                raise ImportError_(f"Linha {n}: falta o ticker.")
            instrument = _instrument(ticker, name_, item["currency"], isin, symbol, kind)
        db.session.add(Transaction(broker=broker, instrument=instrument, external_id=fp,
                                   import_batch=batch.id, **item))
        batch.imported += 1
    _store_observations(fx_obs, price_obs)
    batch.ignored = ignored
    if dates:
        batch.first_date, batch.last_date = min(dates), max(dates)
    db.session.commit()
    return batch


def _store_observations(fx_obs, price_obs):
    """Guarda câmbios e preços observados no extrato quando ainda não há valor para esse dia.

    Servem de base aos gráficos sem ligação à internet; a atualização de mercado substitui-os
    pelas cotações de fecho reais.
    """
    if fx_obs:
        have = {(c, d) for c, d in db.session.query(FxRate.currency, FxRate.date)
                .filter(FxRate.currency.in_({c for c, _ in fx_obs}))}
        for (cur, d), rates in fx_obs.items():
            if (cur, d) not in have:
                db.session.add(FxRate(currency=cur, date=d, rate=sum(rates) / len(rates)))
    if price_obs:
        have = {(i, d) for i, d in db.session.query(Price.instrument, Price.date)
                .filter(Price.instrument.in_({i for i, _ in price_obs}))}
        currencies = {i.ticker: i.currency for i in Instrument.query.filter(
            Instrument.ticker.in_({i for i, _ in price_obs}))}
        for (inst, d), (price, cur) in price_obs.items():
            price = _convert_price(price, cur, currencies.get(inst, cur), d)
            if (inst, d) not in have and price and price > 0:
                db.session.add(Price(instrument=inst, date=d, close=price))


def _fx_on(currency, d):
    """Unidades de `currency` por 1 EUR na data (ou na mais próxima anterior/seguinte)."""
    if currency == "EUR":
        return 1.0
    row = (FxRate.query.filter(FxRate.currency == currency, FxRate.date <= d).order_by(FxRate.date.desc()).first()
           or FxRate.query.filter(FxRate.currency == currency).order_by(FxRate.date).first())
    return row.rate if row else None


def _convert_price(price, from_cur, to_cur, d):
    """Converte um preço observado para a moeda do ativo (ex.: compra em EUR de uma ação cotada em USD)."""
    if from_cur == to_cur:
        return price
    a, b = _fx_on(from_cur, d), _fx_on(to_cur, d)
    return price / a * b if a and b else None


def export_csv(rows, instruments, brokers):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["date", "type", "ticker", "quantity", "price", "amount", "fees", "tax", "currency", "broker", "name"])
    for t in rows:
        inst = instruments.get(t.instrument)
        w.writerow([t.date, t.type, t.instrument or "", t.quantity or "", t.price or "", t.amount or "",
                    t.fees or "", t.tax or "", t.currency, brokers.get(t.broker, ""), inst.name if inst else ""])
    return out.getvalue()
