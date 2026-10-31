"""Testes da aplicação Flask (base de dados SQLite temporária)."""
import io
import re
from pathlib import Path
from datetime import date, timedelta

import pytest

from portfolio_investimentos import create_app
from portfolio_investimentos.models import ImportBatch, Transaction, db

PAGES = ["/", "/carteira", "/desempenho", "/desempenho?periodo=1A&bench=%5EGSPC", "/dividendos", "/fiscal",
         "/radar", "/importar", "/transacoes", "/transacoes?tipo=BUY", "/definicoes"]

T212 = """Action,Time,ISIN,Ticker,Name,No. of shares,Price / share,Currency (Price / share),Exchange rate,Total,Currency (Total),Withholding tax,Currency (Withholding tax),ID
Deposit,2024-01-02 10:00:00,,,,,,,,1000.00,EUR,,,DEP1
Market buy,2024-01-03 15:30:00,US0378331005,AAPL,Apple,2,185.00,USD,1.0950,338.43,EUR,,,EOF1
Dividend (Ordinary),2024-02-15 12:00:00,US0378331005,AAPL,Apple,2,0.24,USD,1.08,0.38,EUR,0.07,USD,
Interest on cash,2024-03-01 00:00:00,,,,,,,,1.25,EUR,,,INT1
"""


@pytest.fixture()
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("PORTFOLIO_DATA_DIR", str(tmp_path))
    app = create_app({"TESTING": True})
    yield app


@pytest.fixture()
def client(app):
    return app.test_client()


def csrf(client):
    html = client.get("/").get_data(as_text=True)
    return re.search(r'csrf_token = "([0-9a-f]+)"', html).group(1)


def post(client, action, token, **data):
    return client.post(f"/api/{action}", data=data, headers={"X-CSRF-Token": token})


def test_empty_pages_render(client):
    for url in PAGES:
        assert client.get(url).status_code == 200, url


def test_demo_pages_render(client):
    token = csrf(client)
    assert post(client, "load_demo", token).status_code == 200
    for url in PAGES:
        resp = client.get(url)
        assert resp.status_code == 200, url
        assert "Traceback" not in resp.get_data(as_text=True)


def test_csrf_required(client):
    assert client.post("/api/hide_banner").status_code == 403
    assert post(client, "hide_banner", "errado").status_code == 403


def test_import_deduplicates(app, client):
    token = csrf(client)
    for expected in ((4, 0), (0, 4)):
        resp = client.post("/api/upload_statement", headers={"X-CSRF-Token": token},
                           data={"broker_key": "trading212", "file": (io.BytesIO(T212.encode()), "t212.csv")},
                           content_type="multipart/form-data")
        assert resp.status_code == 200, resp.get_data(as_text=True)
        r = resp.get_json()["message"][0]
        assert (r["imported"], r["skipped"]) == expected
    with app.app_context():
        assert Transaction.query.count() == 4
        assert ImportBatch.query.count() == 2
        div = Transaction.query.filter_by(type="DIVIDEND").one()
        assert (div.currency, div.tax) == ("USD", pytest.approx(0.07))


def test_broker_selector_counts_open_positions_like_portfolio(app, client):
    """O seletor de corretoras conta as posições como a página Carteira (transferências incluídas)."""
    from portfolio_investimentos.services import data

    token = csrf(client)
    bid = post(client, "save_broker", token, broker_name="Teste").get_json()["message"]
    for typ, qty in (("BUY", "10"), ("BUY", "5"), ("TRANSFER_OUT", "10")):
        inst = "AAA" if qty == "10" else "BBB"
        assert post(client, "add_transaction", token, date="2024-03-01", type=typ, instrument=inst, quantity=qty,
                    price="10", currency="EUR", broker=str(bid)).status_code == 200
    with app.test_request_context():
        b = next(x for x in data.brokers() if x.id == bid)
        assert b.positions == len(data.load_portfolio(str(bid)).open_positions) == 1  # AAA saiu para outra corretora


def test_generic_csv_keeps_identical_rows_same_day(app, client):
    """Duas ordens iguais no mesmo dia são duas transações; reimportar o ficheiro continua a ser duplicado."""
    csv = (b"date,type,ticker,quantity,price,amount,fees,tax,currency,name\n"
           b"2024-01-02,DEPOSIT,,,,2000,,,EUR,\n"
           b"2024-01-03,BUY,ACME,10,50,,,,EUR,Acme\n"
           b"2024-01-03,BUY,ACME,10,50,,,,EUR,Acme\n")
    token = csrf(client)

    def upload():
        resp = client.post("/api/upload_statement", headers={"X-CSRF-Token": token},
                           data={"broker_key": "generic", "file": (io.BytesIO(csv), "a.csv")},
                           content_type="multipart/form-data")
        return resp.get_json()["message"][0]

    first = upload()
    assert (first["imported"], first["skipped"]) == (3, 0)
    again = upload()
    assert (again["imported"], again["skipped"]) == (0, 3)
    with app.app_context():
        assert sum(t.quantity for t in Transaction.query.filter_by(type="BUY")) == 20


def test_api_invalid_input_returns_message_not_500(app, client):
    token = csrf(client)
    first = post(client, "save_broker", token, broker_name="Alfa").get_json()["message"]
    post(client, "save_broker", token, broker_name="Beta")
    dup = post(client, "save_broker", token, broker_name="Beta", name=str(first))
    assert dup.status_code == 400 and "Já existe" in dup.get_json()["error"]
    assert post(client, "save_broker", token, broker_name="Alfa", name=str(first)).status_code == 200  # o próprio nome
    bad = post(client, "performance", token, period="1A", benchmarks="{x")
    assert bad.status_code == 400 and "inválidos" in bad.get_json()["error"]
    assert post(client, "delete_transaction", token).status_code == 400  # falta o argumento


def test_reset_data_clears_refresh_state(app, client):
    from portfolio_investimentos.models import Setting

    with app.app_context():
        for key in ("last_refresh", "refresh_report", "refresh_missing", "hide_banner"):
            Setting.set(key, "1")
        db.session.commit()
    assert post(client, "reset_data", csrf(client), confirm="APAGAR").status_code == 200
    with app.app_context():
        assert all(Setting.get(k) is None for k in ("last_refresh", "refresh_report", "refresh_missing", "hide_banner"))


def test_manual_sell_above_held_quantity_is_rejected(app, client):
    token = csrf(client)
    post(client, "add_transaction", token, date="2024-01-03", type="BUY", instrument="ACME", quantity="5", price="10")
    over = post(client, "add_transaction", token, date="2024-01-04", type="SELL", instrument="ACME", quantity="10",
                price="20")
    assert over.status_code == 400 and "Só tens 5 unidades de ACME" in over.get_json()["error"]
    before = post(client, "add_transaction", token, date="2024-01-02", type="SELL", instrument="ACME", quantity="1",
                  price="20")
    assert before.status_code == 400  # ainda não tinha comprado nessa data
    assert post(client, "add_transaction", token, date="2024-01-04", type="SELL", instrument="ACME", quantity="5",
                price="20").status_code == 200


def test_manual_transfer_and_convert_keep_instrument(app, client):
    """Transferências e conversões manuais têm ativo: antes ficavam sem ele e bloqueavam todas as páginas."""
    token = csrf(client)
    for typ in ("TRANSFER_IN", "TRANSFER_OUT", "CONVERT_IN", "CONVERT_OUT"):
        ok = post(client, "add_transaction", token, date="2024-05-02", type=typ, instrument="acme", quantity="2",
                  price="10", currency="EUR")
        assert ok.status_code == 200, typ
    assert post(client, "add_transaction", token, date="2024-05-02", type="TRANSFER_IN",
                quantity="2").status_code == 400  # sem ticker
    with app.app_context():
        assert Transaction.query.filter(Transaction.instrument.is_(None), Transaction.type != "DEPOSIT").count() == 0
    for page in ("/", "/carteira", "/transacoes", "/definicoes"):
        assert client.get(page).status_code == 200, page


def test_orphan_asset_transaction_does_not_block_pages(app, client):
    """Dado inválido já gravado (transferência sem ativo) é ignorado com aviso em vez de dar erro 500."""
    with app.app_context():
        db.session.add(Transaction(date=date(2024, 1, 2), type="DEPOSIT", amount=1000.0, currency="EUR"))
        db.session.add(Transaction(date=date(2024, 1, 3), type="TRANSFER_IN", instrument=None, quantity=1.0,
                                   price=10.0, currency="EUR"))
        db.session.commit()
    for page in ("/", "/carteira", "/transacoes"):
        assert client.get(page).status_code == 200, page


def test_manual_transaction_and_validation(app, client):
    token = csrf(client)
    ok = post(client, "add_transaction", token, date="2024-05-02", type="BUY", instrument="vwce.de",
              quantity="3", price="110,5", fees="1", currency="eur")
    assert ok.status_code == 200
    bad = post(client, "add_transaction", token, date="2024-05-02", type="BUY", instrument="X", quantity="0")
    assert bad.status_code == 400 and "quantidade" in bad.get_json()["error"]
    with app.app_context():
        t = db.session.get(Transaction, ok.get_json()["message"])
        assert (t.instrument, t.price, t.currency) == ("VWCE.DE", 110.5, "EUR")


def test_backup_and_reset(app, client):
    token = csrf(client)
    post(client, "load_demo", token)
    backup = client.get("/api/backup")
    assert backup.status_code == 200 and backup.data[:15] == b"SQLite format 3"
    import sqlite3 as _sq
    import tempfile as _tf
    with _tf.NamedTemporaryFile(suffix=".db", delete=False) as f:
        f.write(backup.data)
    check = _sq.connect(f.name)
    assert check.execute("PRAGMA integrity_check").fetchone() == ("ok",)  # cópia consistente e abrível
    check.close()
    assert post(client, "reset_data", token, confirm="não").status_code == 400
    assert post(client, "reset_data", token, confirm="APAGAR").status_code == 200
    with app.app_context():
        assert Transaction.query.count() == 0


T212_UTC = """Action,Time (UTC),ISIN,Ticker,Name,Notes,ID,No. of shares,Price / share,Currency (Price / share),Exchange rate,Result,Currency (Result),Total,Currency (Total),Withholding tax,Currency (Withholding tax),Stamp duty reserve tax,Currency (Stamp duty reserve tax),Currency conversion from amount,Currency (Currency conversion from amount),Currency conversion to amount,Currency (Currency conversion to amount),Currency conversion fee,Currency (Currency conversion fee),Merchant name
Deposit,2026-01-02 08:00:00+00:00,,,,Bank Transfer,D1,,,,,,,2000.00,EUR,,,,,,,,,,,
Market buy,2026-01-05 15:00:00+00:00,GB00B63QSB39,GRG,Greggs,,B1,10.0000000000,1500.0000000000,GBX,87.00000000,,,174.91,EUR,,,1.92,EUR,,,,,0.58,EUR,
Market sell,2026-01-06 15:00:00+00:00,US00724F1012,ADBE,Adobe,,S1,3.0000000000,250.0000000000,USD,1.16000000,-40.00,EUR,645.77,EUR,,,,,,,,,0.78,EUR,
Market sell,2026-01-06 15:00:01+00:00,US00724F1012,ADBE,Adobe,,S2,0.0001000000,250.0000000000,USD,1.00000000,0.00,EUR,0.00,EUR,,,,,,,,,,,
Dividend (Dividend),2026-01-07 15:00:00+00:00,US71654V1017,PBR/A,Petrobras,,,40.0000000000,0.2000000,USD,0.85000000,,,6.80,EUR,1.20,USD,,,,,,,,,
Transfer out,2026-01-08 01:00:00+00:00,IE00BK5BQT80,VWCE,Vanguard FTSE All-World (Acc),,T1,3.0000000000,140.0000000000,EUR,1.00000000,0.00,EUR,420.00,EUR,,,,,,,,,,,
Card debit,2026-01-09 19:00:00+00:00,,,,,C1,,,,,,,-100.00,EUR,,,,,,,,,,,Supermercado
Spin off,2026-01-10 09:00:00+00:00,US9256991067,VSNT,Versant Media Group,,SP1,0.5000000000,0E-10,USD,,,,0.00,EUR,,,,,,,,,,,
Currency conversion,2026-01-11 23:00:00+00:00,,,,500.00 EUR -> 580.00 USD,CC1,,,,,,,-0.90,USD,,,,,500.00,EUR,580.00,USD,-0.90,USD,
ADR Fee,2026-01-12 13:00:00+00:00,US31810T1016,FINV,FinVolution,,,300.00000000,-0.02000000,USD,0.86000000,,,-6.00,USD,,,,,,,,,,,
Interest on cash,2026-01-13 02:00:00+00:00,,,,Interest on cash,I1,,,,,,,0.18,EUR,,,,,,,,,,,
"""


def test_import_trading212_full_export(app, client):
    from portfolio_investimentos.models import FxRate, Instrument
    from portfolio_investimentos.services import data

    token = csrf(client)
    resp = client.post("/api/upload_statement", headers={"X-CSRF-Token": token},
                       data={"broker_key": "trading212", "file": (io.BytesIO(T212_UTC.encode()), "t212.csv")},
                       content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    assert resp.get_json()["message"][0]["imported"] == 11
    with app.test_request_context():
        types = sorted(t.type for t in Transaction.query)
        assert types.count("TRANSFER_OUT") == 1 and types.count("FEE") == 2 and types.count("WITHDRAWAL") == 1
        grg = db.session.get(Instrument, "GRG")
        assert grg.currency == "GBP" and grg.data_symbol == "GRG.L"
        buy = Transaction.query.filter_by(instrument="GRG").one()
        assert buy.price == pytest.approx(15.00)
        assert buy.fees == pytest.approx((1.92 + 0.58) * 0.87)
        div = Transaction.query.filter_by(type="DIVIDEND").one()
        assert div.currency == "USD" and div.tax == pytest.approx(1.20)
        assert div.amount == pytest.approx(8.0 + 1.20)
        # a venda residual (total 0, taxa 1.0) não pode estragar o câmbio do dia
        fx = FxRate.query.filter_by(currency="USD").all()
        assert all(abs(f.rate - 1.0) > 0.01 for f in fx)
        p = data.load_portfolio(None)
        # o preço da compra serve de cotação enquanto não há fechos reais
        assert p.book.price("GRG", date(2026, 1, 5)) == pytest.approx(15.00)
        assert {o.asset.ticker for o in p.opening} == {"ADBE", "VWCE"}
        assert p.realized == pytest.approx(-40.0, abs=0.01)


def _xtb_workbook():
    """XLSX sintético com a estrutura do extrato da XTB."""
    from datetime import datetime as dt

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Closed Positions"
    for row in (["Account number", "123"], ["Closed Positions", ""],
                ["Instrument", "Ticker", "Category", "Type", "Volume", "Open Price", "Open Time (UTC)", "Close Price",
                 "Close Time (UTC)", "Product", "Profit/Loss", "Position ID"],
                ["Prime Global", "F50A.DE", "ETF", "BUY", 10.0, 30.0, dt(2026, 4, 22, 11, 0), 30.0,
                 dt(2026, 4, 23, 13, 0), "My Trades", 0.0, 900.0],                       # fecho técnico
                ["Prime Global", "F50A.DE", "ETF", "BUY", 10.0, 30.0, dt(2026, 4, 23, 13, 1), 36.0,
                 dt(2026, 5, 4, 8, 3, 24), "My Trades", 60.0, 901.0],                   # vendido com dinheiro
                ["Profit/loss", 60.0]):
        ws.append(row)
    ops = wb.create_sheet("Cash Operations")
    for row in (["Account number", "123"], ["Cash Operations", ""],
                ["Type", "Instrument", "Ticker", "Category", "Time", "Amount", "ID", "Comment", "Product", "Position ID"],
                ["Deposit", "", "", "", dt(2026, 1, 2, 7), 1000.0, "1", "Deposit", "My Trades", ""],
                ["Stock purchase", "Core S&P 500", "SXR8.DE", "ETF", dt(2026, 1, 3, 9), -480.0, "2",
                 "OPEN BUY 1 @ 480.00", "My Trades", "100"],
                ["Stock purchase", "Ultrashort Bond", "ERNX.DE", "ETF", dt(2026, 1, 5, 5), -55.0, "3",
                 "OPEN BUY ERNX.DE 10/10.0000 @ 5.5000", "Investment Plans", "101"],
                ["Subaccount transfer", "", "", "", dt(2026, 1, 6, 5), -50.0, "4", "Transfer", "My Trades", ""],
                ["Subaccount transfer", "", "", "", dt(2026, 1, 6, 5), 50.0, "5", "Transfer", "Investment Plans", ""],
                ["Stock sell", "Prime Global", "F50A.DE", "ETF", dt(2026, 5, 4, 8, 3, 25), 360.0, "6",
                 "CLOSE BUY 10/10 @ 36.000", "My Trades", "901"],
                ["Free funds interest", "", "", "", dt(2026, 6, 1, 7), 0.20, "7", "Free-funds Interest", "My Trades", ""],
                ["Free funds interest tax", "", "", "", dt(2026, 6, 1, 7), -0.06, "8", "Interest Tax", "My Trades", ""],
                ["Close trade", "", "ETHEREUM", "CRT", dt(2026, 6, 2, 7), -0.38, "9", "Profit of position", "My Trades", "7"],
                ["Total", 775.30]):
        ops.append(row)
    op = wb.create_sheet("Open Positions")
    for row in (["Account number", "123"], ["Open Positions", ""], ["Data as of report generated", dt(2026, 10, 24, 15)],
                ["Product", "Metric", "Amount", "Currency"], ["My Trades", "Open position value", 590.0, "EUR"],
                ["Product", "Instrument/Position", "Ticker", "Category", "Type", "Volume", "Value", "Current price",
                 "Open price", "Open time (UTC)"],
                ["My Trades", "100", "SXR8.DE", "", "BUY", 1.0, 520.0, 520.0, 480.0, dt(2026, 1, 3, 9)],
                ["Investment Plan", "101", "ERNX.DE", "", "BUY", 10.0, 56.0, 5.60, 5.50, dt(2026, 1, 5, 5)]):
        op.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xtb_short_workbook():
    """Extrato da XTB de poucos dias: só um depósito, mas a folha Open Positions lista todas as posições."""
    from datetime import datetime as dt

    import openpyxl
    wb = openpyxl.Workbook()
    ops = wb.active
    ops.title = "Cash Operations"
    for row in (["Account number", "123"], ["Cash Operations", ""], ["Date from (UTC)", dt(2026, 10, 22, 23)],
                ["Date to (UTC)", dt(2026, 10, 26, 7)],
                ["Type", "Instrument", "Ticker", "Category", "Time", "Amount", "ID", "Comment", "Product", "Position ID"],
                ["Deposit", "", "", "", dt(2026, 10, 26, 4), 200.0, "50", "Deposit", "Investment Plans", ""],
                ["Total", None, None, None, None, 200.0]):
        ops.append(row)
    op = wb.create_sheet("Open Positions")
    for row in (["Account number", "123"], ["Open Positions", ""], ["Data as of report generated", dt(2026, 10, 26, 7)],
                ["Product", "Instrument/Position", "Ticker", "Category", "Type", "Volume", "Value", "Current price",
                 "Open price", "Open time (UTC)"],
                ["My Trades", "Core S&P 500", "SXR8.DE", "ETF", "", 1.0, 530.0, "", 480.0, ""],
                ["My Trades", "100", "SXR8.DE", "", "BUY", 1.0, 530.0, 530.0, 480.0, dt(2026, 1, 3, 9)],
                ["Investment Plan", "101", "ERNX.DE", "", "BUY", 10.0, 56.0, 5.60, 5.50, dt(2026, 1, 5, 5)]):
        op.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload_xtb(client, data, name):
    resp = client.post("/api/upload_statement", headers={"X-CSRF-Token": csrf(client)},
                       data={"broker_key": "xtb", "file": (io.BytesIO(data), name)},
                       content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    return resp.get_json()["message"][0]


def test_import_xtb_short_period_does_not_duplicate(app, client):
    """Extrato de poucos dias por cima do histórico: as posições compradas antes não são transferências."""
    from portfolio_investimentos.services import data

    _upload_xtb(client, _xtb_workbook(), "historico.xlsx")
    r = _upload_xtb(client, _xtb_short_workbook(), "semana.xlsx")
    assert r["imported"] == 1  # só o depósito
    assert r["ignored"] == {"Posições compradas antes do período do extrato (já importadas)": 2}
    with app.test_request_context():
        assert Transaction.query.filter_by(type="TRANSFER_IN").count() == 1  # a do histórico
        assert data.load_portfolio(None).holdings_value == pytest.approx(530.0 + 56.0)


def test_import_xtb_short_period_first_import(app, client):
    """Sem histórico, o extrato curto é a única fonte: as posições abertas entram como transferências."""
    r = _upload_xtb(client, _xtb_short_workbook(), "semana.xlsx")
    assert r["imported"] == 3
    with app.test_request_context():
        assert Transaction.query.filter_by(type="TRANSFER_IN").count() == 2


def test_import_xtb_xlsx(app, client):
    from portfolio_investimentos.models import Instrument, Price
    from portfolio_investimentos.services import data

    token = csrf(client)
    resp = client.post("/api/upload_statement", headers={"X-CSRF-Token": token},
                       data={"broker_key": "xtb", "file": (io.BytesIO(_xtb_workbook()), "xtb.xlsx")},
                       content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    r = resp.get_json()["message"][0]
    assert r["ignored"] == {"Transferência entre subcontas": 2}
    with app.test_request_context():
        assert db.session.get(Instrument, "SXR8.DE").data_symbol == "SXR8.DE"
        assert db.session.get(Instrument, "F50A.DE").kind == "etf"
        transfer = Transaction.query.filter_by(type="TRANSFER_IN").one()  # só a parte vendida com dinheiro
        assert (transfer.quantity, transfer.price) == (10.0, 30.0)
        # preço atual das posições abertas usado como cotação na data do relatório
        assert Price.query.filter_by(instrument="SXR8.DE").order_by(Price.date.desc()).first().close == 520.0
        p = data.load_portfolio(None)
        assert p.warnings == []
        assert p.realized == pytest.approx(60.0)
        assert p.holdings_value == pytest.approx(520.0 + 56.0)
        # caixa: 1000 - 480 - 55 + 360 + 0.20 - 0.06 - 0.38
        assert p.cash == pytest.approx(824.76)


TR_CSV = '''"datetime","date","account_type","category","type","asset_class","name","symbol","shares","price","amount","fee","tax","currency","original_amount","original_currency","fx_rate","description","transaction_id","counterparty_name","counterparty_iban","payment_reference","mcc_code"
"2026-09-20T10:00:00Z","2026-09-20","DEFAULT","CASH","TRANSFER_INSTANT_INBOUND","","Eu","","","","1500.000000","","","EUR","","","","Incoming transfer","t1","","","",""
"2026-09-23T07:28:00Z","2026-09-23","DEFAULT","TRADING","BUY","STOCK","FinVolution (ADR)","US31810T1016","80.0000000000","3.1000000000","-248.00","-1.00","-0.04","EUR","","","","Buy trade","t2","","","",""
"2026-09-24T15:09:00Z","2026-09-24","DEFAULT","TRADING","BUY","STOCK","LVMH","FR0000121014","1.0000000000","398.8000000000","-398.80","-1.00","-1.64","EUR","","","","Buy trade","t3","","","",""
"2026-09-25T15:09:00Z","2026-09-25","DEFAULT","TRADING","SELL","STOCK","LVMH","FR0000121014","-1.0000000000","400.0000000000","400.00","-1.00","-0.04","EUR","","","","Sell trade","t4","","","",""
"2026-10-01T08:00:00Z","2026-10-01","DEFAULT","CASH","INTEREST_PAYMENT","","","","","","1.120000","","","EUR","","","","Interest payment","t5","","","",""
"2026-10-02T08:00:00Z","2026-10-02","DEFAULT","TRADING","BUY","FUND","FTSE All-World","IE000716YHJ7","5.0000000000","9.0000000000","-45.00","","","EUR","","","","Savings plan execution","t6","","","",""
'''


def _upload(client, token, key, name, content):
    return client.post("/api/upload_statement", headers={"X-CSRF-Token": token},
                       data={"broker_key": key, "file": (io.BytesIO(content), name)},
                       content_type="multipart/form-data")


@pytest.mark.parametrize("trade_republic_first", [True, False])
def test_trade_republic_and_isin_unification(app, client, trade_republic_first):
    from portfolio_investimentos.models import Instrument
    from portfolio_investimentos.services import data

    t212 = ("Action,Time (UTC),ISIN,Ticker,Name,ID,No. of shares,Price / share,Currency (Price / share),"
            "Exchange rate,Total,Currency (Total)\n"
            "Deposit,2026-09-01 08:00:00+00:00,,,,D1,,,,,500.00,EUR\n"
            "Market buy,2026-09-18 15:00:00+00:00,US31810T1016,FINV,FinVolution,B1,50,3.50,USD,1.17,149.57,EUR\n")
    token = csrf(client)
    jobs = [("traderepublic", "tr.csv", TR_CSV.encode()), ("trading212", "t212.csv", t212.encode())]
    for key, name, content in (jobs if trade_republic_first else jobs[::-1]):
        assert _upload(client, token, key, name, content).status_code == 200
    with app.test_request_context():
        # o mesmo título nas duas corretoras é um só ativo, com o ticker real da Trading 212
        assert Transaction.query.filter_by(instrument="FINV").count() == 2
        assert db.session.get(Instrument, "US31810T1016") is None
        p = data.load_portfolio(None)
        finv = p.positions["FINV"]
        assert finv.qty == pytest.approx(130)
        # cotação mais recente = compra na Trade Republic (3,10 EUR), convertida para USD
        assert finv.value_eur == pytest.approx(130 * 3.10, rel=0.01)
        assert p.realized == pytest.approx((400 - 1.04) - (398.8 + 2.64))
        assert p.interest_eur == pytest.approx(1.12)


class _Hist:
    """Imitação mínima de um DataFrame do yfinance (só o que a atualização usa)."""

    def __init__(self, rows):
        from datetime import datetime
        self.rows = [(datetime.combine(d, datetime.min.time()), {"Close": c}) for d, c in rows]
        self.empty = not rows

    def iterrows(self):
        return iter(self.rows)


class _FakeYf:
    QUOTES = {  # símbolo -> (moeda, fecho)
        "ADBE": ("USD", 263.82), "ADB.DE": ("EUR", 227.10),  # a cotação dos EUA aparece primeiro na pesquisa
        "NVO": ("USD", 39.05), "NOVO-B.CO": ("DKK", 260.15), "NOV.DE": ("EUR", 34.80),
        "GRG.L": ("GBp", 1650.0),
    }
    SEARCH = {"US00724F1012": ["ADBE", "ADB.DE"], "DK0062498333": ["NVO", "NOVO-B.CO"]}

    class Search:
        def __init__(self, isin, max_results=10):
            self.quotes = [{"symbol": s} for s in _FakeYf.SEARCH.get(isin, [])]

    class Ticker:
        def __init__(self, symbol):
            self.symbol, self.news, self.calendar, self.info = symbol, [], {}, {}
            self.history_metadata = {}

        def history(self, start=None, **kw):
            cur, close = _FakeYf.QUOTES.get(self.symbol, (None, None))
            self.history_metadata = {"currency": cur}
            # o Yahoo devolve às vezes dias sem fecho (NaN)
            return _Hist([(date(2026, 10, 21), float("nan")), (date(2026, 10, 22), close),
                          (date(2026, 10, 23), close)] if cur else [])


def test_refresh_picks_symbol_in_instrument_currency(app, client, monkeypatch):
    import sys
    from portfolio_investimentos.models import Instrument, Price
    from portfolio_investimentos.services import market
    monkeypatch.setitem(sys.modules, "yfinance", _FakeYf)
    with app.app_context():
        db.session.add_all([
            # símbolo errado guardado por uma versão anterior (cotação em USD num ativo em EUR)
            Instrument(ticker="ADB", currency="EUR", isin="US00724F1012", data_symbol="ADBE"),
            Instrument(ticker="NOV", currency="EUR", isin="DK0062498333"),
            Instrument(ticker="GRG", currency="GBP", isin="GB00B63QSB39", data_symbol="GRG.L"),
            Price(instrument="ADB", date=date(2026, 10, 1), close=270.0),
            Transaction(date=date(2026, 8, 27), type="BUY", instrument="ADB", quantity=1, price=240.4, currency="EUR"),
            Transaction(date=date(2026, 9, 14), type="BUY", instrument="NOV", quantity=1, price=37.5, currency="EUR"),
            Instrument(ticker="XYZ", currency="EUR"),  # não existe no Yahoo Finance
            Transaction(date=date(2026, 9, 14), type="BUY", instrument="XYZ", quantity=1, price=10.0, currency="EUR"),
        ])
        db.session.commit()
    assert "Ainda não atualizaste" in client.get("/").get_data(as_text=True)
    with app.app_context():
        market.refresh_all(start=date(2026, 10, 1))
        assert db.session.get(Instrument, "ADB").data_symbol == "ADB.DE"
        assert db.session.get(Instrument, "NOV").data_symbol == "NOV.DE"
        assert [p.close for p in Price.query.filter_by(instrument="ADB")] == [227.10, 227.10]
        assert [p.close for p in Price.query.filter_by(instrument="GRG")] == pytest.approx([16.5, 16.5])
    html = client.get("/").get_data(as_text=True)
    assert "Ainda não atualizaste" not in html
    assert "Sem cotações do Yahoo Finance para XYZ" in html and "ADB," not in html
    assert "ADB.DE: 2 cotações" in client.get("/definicoes").get_data(as_text=True)
    st = client.post("/api/refresh_status", headers={"X-CSRF-Token": csrf(client)}).get_json()["message"]
    assert (st["running"], st["error"], st["missing"], st["updated"]) == (False, None, ["XYZ"], 3)


def test_refresh_merges_isin_placeholder_into_exchange_ticker(app, monkeypatch):
    """O mesmo ETF na XTB (ticker da bolsa) e na Trade Republic (ISIN) passa a ser um só ativo."""
    import sys
    from portfolio_investimentos.models import Instrument, Price
    from portfolio_investimentos.services import market
    monkeypatch.setitem(sys.modules, "yfinance", _FakeYf)
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="ADB.DE", currency="EUR", kind="etf", data_symbol="ADB.DE"),
            Instrument(ticker="US00724F1012", currency="EUR", isin="US00724F1012"),
            Transaction(date=date(2026, 8, 27), type="BUY", instrument="ADB.DE", quantity=1, price=240.4, currency="EUR"),
            Transaction(date=date(2026, 9, 14), type="BUY", instrument="US00724F1012", quantity=2, price=230.0,
                        currency="EUR"),
            Price(instrument="US00724F1012", date=date(2026, 10, 22), close=1.0),  # data que o outro também tem
        ])
        db.session.commit()
        report = market.refresh_all(start=date(2026, 10, 1))
        assert "ativos juntados" in report
        assert db.session.get(Instrument, "US00724F1012") is None
        merged = db.session.get(Instrument, "ADB.DE")
        assert merged.isin == "US00724F1012"
        assert sorted(t.instrument for t in Transaction.query) == ["ADB.DE", "ADB.DE"]
        assert [p.close for p in Price.query.filter_by(instrument="ADB.DE")] == [227.10, 227.10]
        # uma nova importação da Trade Republic (ISIN como ticker) já cai no ativo unificado
        from portfolio_investimentos.services import importer
        assert importer._instrument("US00724F1012", isin="US00724F1012") == "ADB.DE"


def test_settings_assets_show_attention_rows_and_quote_status(app, client):
    from portfolio_investimentos.models import Instrument, Price
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="OPEN", currency="EUR"), Instrument(ticker="SHUT", currency="EUR"),
            Transaction(date=date(2026, 1, 6), type="BUY", instrument="OPEN", quantity=2, price=10.0, currency="EUR"),
            Transaction(date=date(2026, 1, 6), type="BUY", instrument="SHUT", quantity=1, price=10.0, currency="EUR"),
            Transaction(date=date(2026, 2, 6), type="SELL", instrument="SHUT", quantity=1, price=11.0, currency="EUR"),
            Price(instrument="SHUT", date=date(2026, 2, 6), close=11.0),
        ])
        db.session.commit()
    html = client.get("/definicoes").get_data(as_text=True)
    rows = re.findall(r'data-asset-row data-open="(\d)" data-noq="(\d)" data-ticker="(\w+)"', html)
    # a aberta (sem cotações) vê-se por omissão e está em "sem cotações"; a vendida, com cotações, não está em nenhum
    assert rows == [("1", "1", "OPEN"), ("0", "0", "SHUT")]
    assert "sem cotações" in html and "11,00 EUR · 06/02/2026" in html
    # uma lista de "sem cotações" desatualizada não leva uma posição vendida, com cotações, para o filtro
    with app.app_context():
        from portfolio_investimentos.models import Setting
        Setting.set("refresh_missing", "SHUT,OPEN")
    html = client.get("/definicoes").get_data(as_text=True)
    assert re.findall(r'data-asset-row data-open="(\d)" data-noq="(\d)" data-ticker="(\w+)"', html) == [
        ("1", "1", "OPEN"), ("0", "0", "SHUT")]


def test_market_status_holidays_and_us_hours():
    from datetime import datetime
    from portfolio_investimentos.services import market

    L = market.LISBON
    at = lambda *a: datetime(*a, tzinfo=L)  # noqa: E731
    # Páscoa 2026: 5 de abril; Sexta-feira Santa (3) e Segunda-feira de Páscoa (6) fecham a Euronext
    assert market._easter(2026) == date(2026, 4, 5) and market._easter(2027) == date(2027, 3, 28)
    s = market.market_status(at(2026, 4, 3, 10, 0))
    assert (s["open"], s["detail"]) == (False, "Abre terça-feira às 08:00")
    assert market.market_status(at(2026, 12, 25, 10, 0))["open"] is False  # Natal
    assert market.market_status(at(2026, 10, 29, 10, 0))["open"] is True  # dia normal
    # 4 de julho de 2026 é sábado: a NYSE fecha na sexta, dia 3
    assert market._session("US", date(2026, 7, 3)) is None and market._session("US", date(2026, 7, 2))
    assert date(2026, 11, 26) in market.us_holidays(2026)  # Ação de Graças
    # Europa fechada, EUA abertos (18:00 em Lisboa em fevereiro; Nova Iorque fecha às 21:00)
    both = market.market_status(at(2026, 2, 3, 18, 0), ("EU", "US"))
    assert both["open"] is True and both["label"] == "Europa: fechado · EUA: aberto até às 21:00"
    # Nova Iorque ainda em hora de verão, Lisboa já não: a sessão acaba às 20:00 em Lisboa
    assert "EUA: aberto até às 20:00" in market.market_status(at(2026, 10, 29, 18, 0), ("EU", "US"))["label"]
    # tudo fechado: diz quando abre a primeira bolsa
    shut = market.market_status(at(2026, 12, 25, 12, 0), ("EU", "US"))
    assert shut["open"] is False and shut["detail"] == "Abre segunda-feira às 08:00 (Europa)"


def test_market_regions_from_positions(app, client):
    from portfolio_investimentos.models import Instrument
    from portfolio_investimentos.services import context, data
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="VWCE.DE", currency="EUR"), Instrument(ticker="AAPL", currency="USD"),
            Instrument(ticker="IE00B4L5Y983", currency="EUR"),
            Transaction(date=date(2026, 1, 6), type="BUY", instrument="VWCE.DE", quantity=1, price=10.0, currency="EUR"),
            Transaction(date=date(2026, 1, 6), type="BUY", instrument="IE00B4L5Y983", quantity=1, price=10.0, currency="EUR"),
        ])
        db.session.commit()
    with app.test_request_context():
        assert context._regions(data.load_portfolio(None)) == {"EU"}  # o ISIN não conta como ação dos EUA
    with app.app_context():
        db.session.add(Transaction(date=date(2026, 1, 7), type="BUY", instrument="AAPL", quantity=1, price=10.0,
                                   currency="USD"))
        db.session.commit()
    with app.test_request_context():
        assert context._regions(data.load_portfolio(None)) == {"EU", "US"}


def test_overview_chart_payload_is_slim(app, client):
    from datetime import timedelta
    from portfolio_investimentos.views import _sample

    days = [(date(2024, 1, 1) + timedelta(days=i)).isoformat() for i in range(0, 800)]
    weekly = _sample(days, "MAX")
    assert weekly[0] == 0 and weekly[-1] == len(days) - 1 and len(weekly) < 130
    assert _sample(days, "1A") == list(range(len(days)))  # até 1 ano, todos os dias
    token = csrf(client)
    post(client, "load_demo", token)
    full = post(client, "performance", token, period="MAX").get_json()["message"]
    assert set(full) == {"start", "dates", "gain", "value", "benchmarks"}  # nada do que a página não desenha
    assert all(set(b) == {"ticker", "name", "color", "twr"} for b in full["benchmarks"])
    assert len(full["dates"]) < 400 and len(full["dates"]) == len(full["value"]) == len(full["gain"])
    assert all(round(v, 2) == v for v in full["value"])
    assert len(client.get("/").data) < 120_000  # antes só o gráfico pesava mais de 300 KB


def test_fiscal_groups_by_asset_with_lots_on_demand_and_full_csv(app, client):
    from datetime import timedelta
    from portfolio_investimentos.models import Instrument
    with app.app_context():
        db.session.add_all([Instrument(ticker="AAA", currency="EUR"), Instrument(ticker="BBB", currency="EUR")])
        db.session.add(Transaction(date=date(2024, 1, 2), type="DEPOSIT", amount=100000.0, currency="EUR"))
        for i in range(40):  # compras fracionadas: cada uma é um lote FIFO na venda
            d = date(2024, 1, 3) + timedelta(days=i)
            db.session.add(Transaction(date=d, type="BUY", instrument="AAA", quantity=1.0, price=10.0, currency="EUR"))
        db.session.add(Transaction(date=date(2024, 3, 1), type="BUY", instrument="BBB", quantity=2.0, price=50.0,
                                   currency="EUR"))
        db.session.add(Transaction(date=date(2024, 6, 1), type="SELL", instrument="AAA", quantity=40.0, price=12.0,
                                   currency="EUR"))
        db.session.add(Transaction(date=date(2024, 6, 1), type="SELL", instrument="BBB", quantity=2.0, price=55.0,
                                   currency="EUR"))
        db.session.commit()
    html = client.get("/fiscal?ano=2024").get_data(as_text=True)
    assert "2 ativos · 41 linhas FIFO" in html
    assert html.count("data-lots=") == 2 and html.count("<tr>") < 10  # uma linha por ativo, não por lote
    frag = client.get("/fiscal/lotes?ano=2024&ativo=AAA").get_data(as_text=True)
    assert frag.count("<tr>") == 41  # cabeçalho + 40 lotes
    assert client.get("/fiscal/lotes?ano=2024&ativo=NAO").status_code == 404
    assert client.get("/fiscal/lotes?ano=1999&ativo=AAA").status_code == 404
    csv = client.get("/fiscal.csv?ano=2024")
    assert csv.status_code == 200 and "attachment" in csv.headers["Content-Disposition"]
    lines = csv.get_data(as_text=True).lstrip("\ufeff").splitlines()
    assert len(lines) == 1 + 41 and lines[0].startswith("Data venda;Ativo;")
    assert lines[1].split(";")[:2] == ["2024-06-01", '"AAA"'] and lines[1].split(";")[4] == "12.0"


def test_background_refresh_retries_when_database_is_locked(app, monkeypatch):
    from sqlalchemy.exc import OperationalError
    from portfolio_investimentos.models import Setting
    from portfolio_investimentos.services import market

    assert app.config["SQLALCHEMY_ENGINE_OPTIONS"]["connect_args"]["timeout"] == 30
    calls, sleeps = [], []

    def flaky(**kwargs):
        calls.append(1)
        if len(calls) < 3:
            raise OperationalError("UPDATE", {}, Exception("database is locked"))
        return "ok"
    monkeypatch.setattr(market, "refresh_all", flaky)
    monkeypatch.setattr(market._time, "sleep", sleeps.append)

    class Inline:  # corre a thread já, para o teste esperar pelo resultado
        def __init__(self, target, daemon=None):
            self.target = target

        def start(self):
            self.target()
    monkeypatch.setattr(market.threading, "Thread", Inline)
    with app.app_context():
        market.refresh_in_background(app)
        assert len(calls) == 3 and sleeps == [15, 15] and Setting.get("refresh_status") == ""
    # um erro que não é de bloqueio não se repete
    calls.clear()

    def broken(**kwargs):
        calls.append(1)
        raise RuntimeError("falha qualquer")
    monkeypatch.setattr(market, "refresh_all", broken)
    with app.app_context():
        market.refresh_in_background(app)
        assert len(calls) == 1 and Setting.get("refresh_status") == "error:falha qualquer"


def test_merge_instruments_manually(app, client):
    from portfolio_investimentos.models import Instrument, Price
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="AAA.DE", currency="EUR"), Instrument(ticker="AAA", currency="EUR", sector="Tec"),
            Instrument(ticker="USD1", currency="USD"),
            Transaction(date=date(2026, 1, 5), type="BUY", instrument="AAA", quantity=1, price=10.0, currency="EUR"),
            Price(instrument="AAA", date=date(2026, 1, 6), close=11.0),
            Price(instrument="AAA.DE", date=date(2026, 1, 6), close=11.5),
        ])
        db.session.commit()
    token = csrf(client)
    assert post(client, "merge_instruments", token, source="AAA", target="AAA").status_code == 400
    assert post(client, "merge_instruments", token, source="AAA", target="USD1").status_code == 400  # moedas diferentes
    assert post(client, "merge_instruments", token, source="AAA", target="NAO").status_code == 400
    assert post(client, "merge_instruments", token, source="AAA", target="AAA.DE").status_code == 200
    with app.app_context():
        assert db.session.get(Instrument, "AAA") is None
        assert db.session.get(Instrument, "AAA.DE").sector == "Tec"
        assert Transaction.query.filter_by(instrument="AAA.DE").count() == 1
        assert [p.close for p in Price.query.filter_by(instrument="AAA.DE")] == [11.5]
    assert client.get("/definicoes").status_code == 200 and client.get("/carteira").status_code == 200


def test_yfinance_deprecation_warnings_are_silenced(monkeypatch):
    import sys
    import warnings
    from portfolio_investimentos.services import market
    monkeypatch.setitem(sys.modules, "yfinance", _FakeYf)
    with warnings.catch_warnings(record=True) as seen:
        market._yahoo_client()
        warnings.warn_explicit("inplace obsoleto", DeprecationWarning, "yfinance/utils.py", 568, module="yfinance.utils")
        warnings.warn_explicit("outro aviso", DeprecationWarning, "app.py", 1, module="portfolio_investimentos.x")
    assert [str(w.message) for w in seen] == ["outro aviso"]  # só se cala o do yfinance


def test_refresh_only_open_positions(app, monkeypatch):
    """Com only_open, as posições fechadas não se pedem ao Yahoo (nem os seus câmbios)."""
    import sys
    from portfolio_investimentos.models import Instrument
    from portfolio_investimentos.services import market

    calls = []

    class Yf(_FakeYf):
        class Ticker(_FakeYf.Ticker):
            def history(self, start=None, **kw):
                calls.append(self.symbol)
                return super().history(start=start, **kw)

    monkeypatch.setitem(sys.modules, "yfinance", Yf)
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="ADB", currency="EUR", isin="US00724F1012", data_symbol="ADB.DE"),
            Instrument(ticker="NOV", currency="EUR", isin="DK0062498333", data_symbol="NOV.DE"),
            Transaction(date=date(2026, 8, 27), type="BUY", instrument="ADB", quantity=1, price=240.4, currency="EUR"),
            Transaction(date=date(2026, 9, 14), type="BUY", instrument="NOV", quantity=1, price=37.5, currency="EUR"),
            Transaction(date=date(2026, 9, 20), type="SELL", instrument="NOV", quantity=1, price=38.0, currency="EUR"),
        ])
        db.session.commit()
        lines = []
        report = market.refresh_all(only_open=True, progress=lines.append)
        assert "ADB.DE" in calls and "NOV.DE" not in calls  # NOV está vendida
        # o contador só conta o que se atualiza: ADB e os 3 índices de comparação (não a posição vendida)
        assert any(ln.startswith("[1/4] ") for ln in lines) and not any("/5]" in ln for ln in lines)
        assert any(ln.startswith("[1/4] ") and ln.endswith(" s)") for ln in lines)  # com o tempo de cada ativo
        assert any(ln.startswith("[eventos 1/1] ADB") for ln in lines)
        assert "1 ativos fechados não foram atualizados" in report
        calls.clear()
        market.refresh_all()  # sem a opção, pede também a fechada
        assert "NOV.DE" in calls


def test_refresh_reports_progress_with_last_close(app, monkeypatch):
    """A atualização conta os ativos à medida que avança e mostra o último fecho de cada um."""
    import sys
    from portfolio_investimentos.models import Instrument
    from portfolio_investimentos.services import market
    monkeypatch.setitem(sys.modules, "yfinance", _FakeYf)
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="ADB", currency="EUR", isin="US00724F1012", data_symbol="ADB.DE"),
            Instrument(ticker="NOV", currency="EUR", isin="DK0062498333"),
            Transaction(date=date(2026, 8, 27), type="BUY", instrument="ADB", quantity=1, price=240.4, currency="EUR"),
            Transaction(date=date(2026, 9, 14), type="BUY", instrument="NOV", quantity=1, price=37.5, currency="EUR"),
        ])
        db.session.commit()
        lines = []
        report = market.refresh_all(start=date(2026, 10, 1), progress=lines.append)
        asset_lines = [ln for ln in lines if ln.startswith("[")]
        assert any(ln.startswith("[") and "ADB.DE: 2 cotações · último fecho 227.10 EUR (23/10/2026)" in ln
                   for ln in asset_lines)
        assert all("/" in ln.split("]")[0] for ln in asset_lines)  # "[n/total]"
        assert "ADB.DE: 2 cotações · último fecho 227.10 EUR (23/10/2026)" in report  # o resumo guardado não leva o contador
        assert market._progress == {"done": 0, "total": 0, "current": ""}  # reposto no fim


def test_refresh_is_incremental_after_first_run(app, monkeypatch):
    """Depois da primeira atualização, só se pedem os dias novos; posições fechadas já completas não se pedem."""
    import sys
    from portfolio_investimentos.models import Instrument, Price
    from portfolio_investimentos.services import market

    calls = []

    class Yf(_FakeYf):
        class Ticker(_FakeYf.Ticker):
            def history(self, start=None, **kw):
                calls.append((self.symbol, start))
                return super().history(start=start, **kw)

    monkeypatch.setitem(sys.modules, "yfinance", Yf)
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="ADB", currency="EUR", isin="US00724F1012", data_symbol="ADB.DE"),
            Instrument(ticker="NOV", currency="EUR", isin="DK0062498333", data_symbol="NOV.DE"),
            Transaction(date=date(2026, 8, 27), type="BUY", instrument="ADB", quantity=1, price=240.4, currency="EUR"),
            Transaction(date=date(2026, 9, 14), type="BUY", instrument="NOV", quantity=1, price=37.5, currency="EUR"),
            Transaction(date=date(2026, 9, 20), type="SELL", instrument="NOV", quantity=1, price=38.0, currency="EUR"),
        ])
        db.session.commit()
        market.refresh_all(start=date(2026, 8, 1))  # primeira vez: histórico todo
        # o Yahoo simulado só devolve 3 dias; juntam-se cotações antigas, como as de um histórico real
        db.session.add_all([Price(instrument="ADB", date=date(2026, 8, 27), close=240.0),
                            Price(instrument="NOV", date=date(2026, 9, 14), close=37.5)])
        db.session.commit()
        assert db.session.get(Instrument, "ADB").last_quote == date(2026, 10, 23)
        calls.clear()
        report = market.refresh_all()  # seguinte: incremental
        asset_calls = [c for c in calls if c[0] in ("ADB.DE", "NOV.DE")]
        # NOV está vendida e as cotações já passam da última transação: nada a pedir; ADB só desde 18/10
        # (o pedido com 400 dias é o dos dividendos por ação, que se faz para as posições abertas)
        quotes = [c for c in asset_calls if c[1] != date.today() - timedelta(days=400)]
        assert quotes == [("ADB.DE", date(2026, 10, 18))]
        assert "1 ativos só com os últimos dias" in report and "1 posições fechadas" in report
        # símbolo alterado à mão: volta a pedir o histórico todo
        calls.clear()
        db.session.get(Instrument, "ADB").last_quote = None
        db.session.commit()
        market.refresh_all()
        assert [c for c in calls if c[0] == "ADB.DE"][0][1] == date(2026, 8, 20)


def test_refresh_state(app):
    from datetime import datetime, timedelta
    from portfolio_investimentos.models import Setting
    from portfolio_investimentos.services.market import refresh_state
    with app.app_context():
        Setting.set("refresh_status", "running:" + datetime.now().isoformat())
        assert refresh_state()[0] is True
        Setting.set("refresh_status", "running:" + (datetime.now() - timedelta(hours=2)).isoformat())
        assert refresh_state()[0] is False  # a aplicação foi fechada a meio
        Setting.set("refresh_status", "running:" + (datetime.now() - timedelta(minutes=2)).isoformat())
        assert refresh_state()[0] is True
        Setting.set("refresh_status", "running:" + (datetime.now() - timedelta(minutes=6)).isoformat())
        assert refresh_state()[0] is False  # sem sinal de vida há mais de 5 minutos: o botão deixa de rodar
        Setting.set("refresh_status", "error:sem rede")
        assert refresh_state()[1] == "sem rede"


def test_refresh_status_is_cleared_even_if_the_database_is_busy_once(app, monkeypatch):
    from portfolio_investimentos.models import Setting
    from portfolio_investimentos.services import market
    monkeypatch.setattr(market._time, "sleep", lambda s: None)
    real, fails = Setting.set, []

    def busy_once(key, value):
        if not fails:
            fails.append(1)
            raise RuntimeError("database is locked")
        return real(key, value)
    with app.app_context():
        real("refresh_status", "running:" + __import__("datetime").datetime.now().isoformat())
        monkeypatch.setattr(Setting, "set", staticmethod(busy_once))
        assert market._set_status("") is True  # repete e consegue
        assert Setting.get("refresh_status") == ""
        market._background.update(on=True, only_open=True)
        market._progress.update(done=2, total=5, current="eventos de VWCE.DE")
        try:
            market._beat()  # sinal de vida renova a hora e deixa o progresso para qualquer processo do servidor
        finally:
            market._background.update(on=False, only_open=False)
            market._progress.update(done=0, total=0, current="")
        assert Setting.get("refresh_status").startswith("running:")
        assert market.refresh_state()[0] is True
        assert market.refresh_progress() == {"only_open": True, "done": 2, "total": 5, "current": "eventos de VWCE.DE"}
        Setting.set("refresh_status", "")
        assert market.refresh_progress()["total"] == 0


REMOTE = {"REMOTE_ADDR": "203.0.113.7"}


def test_remote_access_blocked_without_password(client):
    assert client.get("/").status_code == 200  # no próprio computador abre sem palavra-passe
    resp = client.get("/", environ_base=REMOTE)
    assert resp.status_code == 403 and "Define uma palavra-passe" in resp.get_data(as_text=True)
    # atrás de um proxy no mesmo servidor o endereço é local, mas o pedido vem de fora
    assert client.get("/", headers={"X-Forwarded-For": "203.0.113.7"}).status_code == 403
    assert client.get("/api/backup", environ_base=REMOTE).status_code == 403


def test_password_login_flow(app, client):
    token = csrf(client)
    assert post(client, "set_password", token, new_password="curta", confirm="curta").status_code == 400
    resp = post(client, "set_password", token, new_password="segredo123", confirm="segredo123")
    assert resp.status_code == 200
    assert client.get("/").status_code == 200  # quem definiu continua com a sessão aberta

    other = app.test_client()  # outro dispositivo
    resp = other.get("/carteira", environ_base=REMOTE)
    assert resp.status_code == 302 and "/login?next=/carteira" in resp.headers["Location"]
    assert other.post("/api/hide_banner", environ_base=REMOTE).status_code == 401
    assert other.get("/api/backup", environ_base=REMOTE).status_code == 401
    login_page = other.get("/login", environ_base=REMOTE).get_data(as_text=True)
    t = re.search(r'name="csrf_token" value="([0-9a-f]+)"', login_page).group(1)
    assert other.post("/login", data={"password": "segredo123"}, environ_base=REMOTE).status_code == 403  # sem CSRF
    bad = other.post("/login", data={"csrf_token": t, "password": "errada"}, environ_base=REMOTE)
    assert bad.status_code == 401 and "Palavra-passe errada" in bad.get_data(as_text=True)
    ok = other.post("/login?next=/%5Cevil.example/x", data={"csrf_token": t, "password": "segredo123"},
                    environ_base=REMOTE)
    assert ok.status_code == 302 and ok.headers["Location"] == "/"  # não redireciona para outro site
    assert other.get("/carteira", environ_base=REMOTE).status_code == 200

    # mudar a palavra-passe termina as sessões dos outros dispositivos
    assert post(client, "set_password", csrf(client), current_password="segredo123",
                new_password="outra-segredo", confirm="outra-segredo").status_code == 200
    assert other.get("/carteira", environ_base=REMOTE).status_code == 302

    # remover a palavra-passe só no próprio computador
    token = csrf(client)
    assert post(client, "remove_password", token, current_password="errada").status_code == 400
    assert post(client, "remove_password", token, current_password="outra-segredo").status_code == 200
    assert client.get("/login").status_code == 302


def test_login_rate_limited(app):
    from portfolio_investimentos import auth
    from portfolio_investimentos.models import Setting
    with app.app_context():
        Setting.set("password_hash", __import__("werkzeug.security").security.generate_password_hash("segredo123"))
    c = app.test_client()
    page = c.get("/login", environ_base=REMOTE).get_data(as_text=True)
    t = re.search(r'name="csrf_token" value="([0-9a-f]+)"', page).group(1)
    for _ in range(auth.MAX_FAILS):
        c.post("/login", data={"csrf_token": t, "password": "x"}, environ_base=REMOTE)
    resp = c.post("/login", data={"csrf_token": t, "password": "segredo123"}, environ_base=REMOTE)
    assert "Demasiadas tentativas" in resp.get_data(as_text=True)
    auth._fails.clear()


def test_env_password(app, monkeypatch):
    monkeypatch.setenv("PORTFOLIO_PASSWORD", "da-variavel")
    c = app.test_client()
    page = c.get("/login", environ_base=REMOTE).get_data(as_text=True)
    t = re.search(r'name="csrf_token" value="([0-9a-f]+)"', page).group(1)
    ok = c.post("/login", data={"csrf_token": t, "password": "da-variavel"}, environ_base=REMOTE)
    assert ok.status_code == 302
    assert "PORTFOLIO_PASSWORD" in c.get("/definicoes", environ_base=REMOTE).get_data(as_text=True)


def test_closes_undo_split_adjustment():
    from portfolio_investimentos.services.market import _closes
    # reagrupamento 1:15 a 07/07/2025: o Yahoo mostra os preços antigos multiplicados por 15
    hist = _Hist([(date(2025, 7, 3), 9.33), (date(2025, 7, 4), float("nan")), (date(2025, 7, 7), 9.40)])
    hist.rows[2][1]["Stock Splits"] = 1 / 15
    assert _closes(hist) == [(date(2025, 7, 3), pytest.approx(0.622)), (date(2025, 7, 7), 9.40)]
    # desdobramento 10:1: os preços antigos aparecem divididos por 10
    hist = _Hist([(date(2024, 6, 7), 120.0), (date(2024, 6, 10), 121.0)])
    hist.rows[1][1]["Stock Splits"] = 10.0
    assert _closes(hist) == [(date(2024, 6, 7), 1200.0), (date(2024, 6, 10), 121.0)]


def test_pwa_files(app, client):
    html = client.get("/").get_data(as_text=True)
    assert 'rel="manifest"' in html
    m = client.get("/static/manifest.webmanifest")
    assert m.status_code == 200 and m.mimetype == "application/manifest+json"
    assert m.get_json()["display"] == "standalone"
    sw = client.get("/sw.js")
    assert sw.status_code == 200 and "javascript" in sw.mimetype and sw.headers["Cache-Control"] == "no-cache"
    for icon in m.get_json()["icons"]:
        assert client.get(icon["src"]).status_code == 200
    # com palavra-passe, o browser continua a conseguir instalar (manifesto e service worker abertos)
    from portfolio_investimentos import auth
    with app.app_context():
        auth.clear_password()
        from portfolio_investimentos.models import Setting
        Setting.set("password_hash", "x")
    other = app.test_client()
    assert other.get("/sw.js", environ_base=REMOTE).status_code == 200
    assert other.get("/static/manifest.webmanifest", environ_base=REMOTE).status_code == 200
    assert other.get("/", environ_base=REMOTE).status_code == 302


def test_portfolio_cache_invalidated_on_changes(app, client):
    from portfolio_investimentos.models import Price
    from portfolio_investimentos.services import data
    token = csrf(client)
    post(client, "add_transaction", token, date="2024-01-02", type="DEPOSIT", amount="1000", currency="EUR")
    post(client, "add_transaction", token, date="2024-01-03", type="BUY", instrument="ACME", quantity="10", price="50",
         currency="EUR")
    with app.test_request_context():
        assert data.load_portfolio(None).holdings_value == pytest.approx(500)
    # nova cotação (como numa atualização): o pedido seguinte já a vê
    with app.app_context():
        db.session.add(Price(instrument="ACME", date=date(2024, 1, 4), close=60.0))
        db.session.commit()
    with app.test_request_context():
        assert data.load_portfolio(None).holdings_value == pytest.approx(600)
    # apagar em bloco (query.delete) também invalida
    with app.app_context():
        Price.query.filter_by(instrument="ACME").delete()
        db.session.commit()
    with app.test_request_context():
        assert data.load_portfolio(None).holdings_value == pytest.approx(500)


def test_legacy_data_dir_is_renamed(tmp_path):
    from portfolio_investimentos import _adopt_legacy
    old, new = tmp_path / "PortfolioManager", tmp_path / "PortfolioInvestimentos"
    old.mkdir()
    (old / "portfolio.db").write_text("dados")
    assert _adopt_legacy(old, new) == new
    assert (new / "portfolio.db").read_text() == "dados" and not old.exists()
    # já com a pasta nova: não mexe em nada
    old.mkdir()
    assert _adopt_legacy(old, new) == new and old.exists()


def test_starts_without_timezone_database(monkeypatch, tmp_path):
    """No Android/Termux sem tzdata, ZoneInfo falha: a aplicação tem de arrancar na mesma."""
    import importlib
    import zoneinfo
    from portfolio_investimentos.services import market

    def no_tz(key):
        raise zoneinfo.ZoneInfoNotFoundError(key)
    monkeypatch.setattr(zoneinfo, "ZoneInfo", no_tz)
    try:
        importlib.reload(market)
        assert isinstance(market.LISBON, market._LisbonFallback)
        assert market.market_status()["label"] in ("Aberto", "Fechado")
        monkeypatch.setenv("PORTFOLIO_DATA_DIR", str(tmp_path))
        app = create_app({"TESTING": True})
        assert app.test_client().get("/").status_code == 200
    finally:
        monkeypatch.undo()
        importlib.reload(market)


def test_detect_broker_and_cli_import(app, tmp_path, monkeypatch, capsys):
    import sys
    from portfolio_investimentos import __main__ as cli
    from portfolio_investimentos.services.importer import detect_broker
    assert detect_broker("x.xlsx", _xtb_workbook()) == "xtb"
    assert detect_broker("t.csv", T212_UTC.encode()) == "trading212"
    assert detect_broker("tr.csv", TR_CSV.encode()) == "traderepublic"
    assert detect_broker("g.csv", b"date,type,ticker,quantity,price\n") == "generic"
    (tmp_path / "xtb.xlsx").write_bytes(_xtb_workbook())
    monkeypatch.setattr(cli, "create_app", lambda: app)
    monkeypatch.setattr(sys, "argv", ["portfolio_investimentos", "--importar", str(tmp_path / "xtb.xlsx"),
                                      str(tmp_path / "nao-existe.csv")])
    with pytest.raises(SystemExit) as exit_:
        cli.main()
    out = capsys.readouterr().out
    assert "xtb.xlsx (xtb): 8 novas transações" in out and "nao-existe.csv: não foi possível ler" in out
    assert exit_.value.code == 1


def test_choose_port(monkeypatch):
    from portfolio_investimentos import __main__ as cli
    busy, ours = {5000, 5001}, {5001}
    monkeypatch.setattr(cli, "_port_in_use", lambda host, port: port in busy)
    monkeypatch.setattr(cli, "_is_this_app", lambda url: any(f":{p}/" in url for p in ours))
    assert cli._choose_port("127.0.0.1", None) == (5001, True)       # a aplicação já está na 5001
    ours.clear()
    assert cli._choose_port("127.0.0.1", None) == (5002, False)      # 5000 e 5001 ocupadas por outros
    assert cli._choose_port("127.0.0.1", 5000) == (None, False)      # porta indicada e ocupada: erro
    assert cli._choose_port("127.0.0.1", 6000) == (6000, False)


# resposta do endpoint "chart" do Yahoo (formato real, valores fictícios)
YAHOO_CHART = {"chart": {"result": [{
    "meta": {"currency": "EUR", "symbol": "ACME.DE", "gmtoffset": 7200},
    "timestamp": [1760943600, 1761030000, 1761116400],          # 20, 21 e 22/10/2025 às 07:00 UTC
    "events": {"splits": {"1761116400": {"date": 1761116400, "numerator": 2, "denominator": 1}},
               "dividends": {"1760943600": {"date": 1760943600, "amount": 0.4}}},
    "indicators": {"quote": [{"close": [50.0, None, 26.0]}]},
}], "error": None}}


def test_refresh_stores_yahoo_dividends_per_share_for_open_positions(app, client, monkeypatch):
    """Ao atualizar, guardam-se os dividendos por ação (Yahoo) das posições abertas; a página usa-os na projeção."""
    import sys
    import time as _t
    from portfolio_investimentos.models import DividendHistory, Instrument
    from portfolio_investimentos.services import market, yahoo
    monkeypatch.setitem(sys.modules, "yfinance", None)  # acesso direto
    now = int(_t.time())
    day = 86400
    chart = {"chart": {"result": [{
        "meta": {"currency": "USD", "gmtoffset": 0},
        "timestamp": [now - 200 * day, now - 110 * day, now - 20 * day, now - 2 * day],
        "events": {"dividends": {str(now - 200 * day): {"amount": 0.5, "date": now - 200 * day},
                                 str(now - 20 * day): {"amount": 0.5, "date": now - 20 * day}}},
        "indicators": {"quote": [{"close": [40.0, 41.0, 42.0, 43.0]}]}}], "error": None}}

    def fake_get(url, params, **kw):
        if "search" in url:
            return {"quotes": [], "news": []}
        return chart if "FINV" in url else {"chart": {"result": None}}
    monkeypatch.setattr(yahoo, "_get_json", fake_get)
    with app.app_context():
        db.session.add_all([
            Instrument(ticker="FINV", currency="USD", data_symbol="FINV"),
            Transaction(date=date.today() - timedelta(days=30), type="DEPOSIT", amount=5000.0, currency="EUR"),
            Transaction(date=date.today() - timedelta(days=30), type="BUY", instrument="FINV", quantity=100, price=40.0,
                        currency="USD"),
            # um dividendo antigo de outro ativo, para a página de dividendos existir
            Instrument(ticker="OLD", currency="EUR"),
            Transaction(date=date.today() - timedelta(days=400), type="BUY", instrument="OLD", quantity=1, price=10.0,
                        currency="EUR"),
            Transaction(date=date.today() - timedelta(days=380), type="DIVIDEND", instrument="OLD", quantity=1, amount=1.0,
                        currency="EUR"),
            Transaction(date=date.today() - timedelta(days=370), type="SELL", instrument="OLD", quantity=1, price=10.0,
                        currency="EUR"),
        ])
        db.session.commit()
        report = market.refresh_all(only_open=True)
        assert DividendHistory.query.filter_by(instrument="FINV").count() == 2
        assert "Dividendos por ação (Yahoo, últimos 12 meses): 1 de 1 posições abertas pagam." in report
    html = client.get("/dividendos").get_data(as_text=True)
    assert "a partir de:" in html and "<b>FINV</b>" in html and "(Yahoo)" in html


def test_yahoo_direct_parse():
    from portfolio_investimentos.services import yahoo
    from portfolio_investimentos.services.market import _closes
    hist, meta = yahoo._parse_chart(YAHOO_CHART)
    assert meta == {"currency": "EUR"} and not hist.empty
    # dia sem fecho ignorado; antes do desdobramento 2:1 o preço negociado era o dobro do ajustado
    assert _closes(hist) == [(date(2025, 10, 20), 100.0), (date(2025, 10, 22), 26.0)]
    assert [r[1]["Dividends"] for r in hist.iterrows()] == [0.4, 0.0, 0.0]  # dividendo no dia ex-dividendo
    assert yahoo._parse_chart({"chart": {"result": None, "error": {"code": "Not Found"}}})[0].empty


def test_refresh_without_yfinance(app, monkeypatch):
    """Sem o yfinance (ex.: Android), a atualização usa o acesso direto ao Yahoo."""
    import sys
    from portfolio_investimentos.models import Instrument, Price
    from portfolio_investimentos.services import market, yahoo
    monkeypatch.setitem(sys.modules, "yfinance", None)  # import yfinance → ImportError
    calls = []

    def fake_get(url, params):
        calls.append(url)
        if "search" in url:
            return {"quotes": [{"symbol": "ACME.DE"}], "news": []}
        return YAHOO_CHART if "ACME.DE" in url else {"chart": {"result": None}}
    monkeypatch.setattr(yahoo, "_get_json", fake_get)
    with app.app_context():
        db.session.add_all([Instrument(ticker="ACME", currency="EUR", isin="DE0000000001"),
                            Transaction(date=date(2025, 10, 1), type="BUY", instrument="ACME", quantity=1,
                                        price=95.0, currency="EUR")])
        db.session.commit()
        assert market.yfinance_available()
        market.refresh_all(start=date(2025, 10, 1))
        assert db.session.get(Instrument, "ACME").data_symbol == "ACME.DE"
        assert [p.close for p in Price.query.filter_by(instrument="ACME").order_by(Price.date)] == [100.0, 26.0]
    assert any("chart/ACME.DE" in u for u in calls)


def test_yahoo_direct_retries_and_crumb(monkeypatch):
    import io as _io
    import json as _json
    import urllib.error
    from portfolio_investimentos.services import yahoo
    monkeypatch.setattr(yahoo, "_state", {"opener": None, "crumb": None})
    monkeypatch.setattr(yahoo._time, "sleep", lambda s: None)
    seen = []

    def http_error(url, code):
        return urllib.error.HTTPError(url, code, "x", {}, _io.BytesIO(b""))

    def fake_fetch(url):
        seen.append(url)
        if "fc.yahoo.com" in url:
            raise http_error(url, 404)
        if "getcrumb" in url:
            return b"abc123"
        if "crumb=abc123" in url:
            return _json.dumps({"ok": True}).encode()
        raise http_error(url, 401 if "query2" in url else 429)
    monkeypatch.setattr(yahoo, "_fetch", fake_fetch)
    assert yahoo._get_json(yahoo.CHART_URL.format(symbol="AAPL"), {"interval": "1d"}) == {"ok": True}
    assert "query1" in seen[0] and "query2" in seen[1] and any("getcrumb" in u for u in seen)

    # sem ligação nenhuma: erro legível, para o resumo da atualização
    monkeypatch.setattr(yahoo, "_state", {"opener": None, "crumb": None})
    monkeypatch.setattr(yahoo, "_fetch", lambda url: (_ for _ in ()).throw(urllib.error.URLError("rede em baixo")))
    with pytest.raises(yahoo.YahooError, match="sem ligação ao Yahoo"):
        yahoo._get_json(yahoo.CHART_URL.format(symbol="AAPL"), {})


def test_refresh_report_shows_yahoo_error(app, monkeypatch):
    import sys
    from portfolio_investimentos.models import Instrument, Setting
    from portfolio_investimentos.services import market, yahoo
    monkeypatch.setitem(sys.modules, "yfinance", None)

    def fail(url, params):
        raise yahoo.YahooError("HTTP 429: o Yahoo limitou os pedidos; tenta mais tarde")
    monkeypatch.setattr(yahoo, "_get_json", fail)
    with app.app_context():
        db.session.add(Instrument(ticker="ACME", currency="EUR", data_symbol="ACME.DE"))
        db.session.commit()
        market.refresh_all(start=date(2025, 10, 1))
        report = Setting.get("refresh_report")
    assert "ACME: sem cotações em EUR (ACME.DE: HTTP 429" in report


def test_offline_mode_marks_pages(app, client, monkeypatch):
    token = csrf(client)
    assert client.get("/ping").status_code == 204
    assert "X-PM-Offline" not in client.get("/").headers          # desligado por omissão
    assert post(client, "set_offline", token, enabled="1").status_code == 200
    page = client.get("/carteira")
    assert page.headers.get("X-PM-Offline") == "1" and 'data-offline="1"' in page.get_data(as_text=True)
    assert "X-PM-Offline" not in client.get("/sw.js").headers
    assert "X-PM-Offline" not in client.get("/static/offline.html").headers
    assert "X-PM-Offline" not in client.get("/definicoes").headers         # só páginas de consulta
    # com palavra-passe: o login e os redirecionamentos nunca ficam guardados
    monkeypatch.setenv("PORTFOLIO_PASSWORD", "segredo123")
    other = app.test_client()
    assert "X-PM-Offline" not in other.get("/login", environ_base=REMOTE).headers
    assert "X-PM-Offline" not in other.get("/", environ_base=REMOTE).headers
    assert other.get("/ping", environ_base=REMOTE).status_code == 204   # aberto, para a página saber se a app corre
    monkeypatch.delenv("PORTFOLIO_PASSWORD")
    assert post(client, "set_offline", token, enabled="0").status_code == 200
    assert "X-PM-Offline" not in client.get("/").headers


# ------------------------------------------------------------ versão do esquema e cópias automáticas
def _legacy_db(folder, monkeypatch):
    """Base de dados de uma versão antiga: sem as colunas `result` e `last_quote` e sem versão do esquema."""
    import sqlite3
    monkeypatch.setenv("PORTFOLIO_DATA_DIR", str(folder))
    app = create_app({"TESTING": True})
    with app.app_context():
        db.session.add(Transaction(date=date(2024, 1, 2), type="DEPOSIT", amount=100.0, currency="EUR"))
        db.session.commit()
        db.session.remove()
        db.engine.dispose()
    conn = sqlite3.connect(folder / "portfolio.db")
    conn.execute('ALTER TABLE "transaction" DROP COLUMN result')
    conn.execute("ALTER TABLE instrument DROP COLUMN last_quote")
    conn.execute("PRAGMA user_version = 0")
    conn.commit()
    conn.close()


def test_new_database_starts_at_latest_schema(app):
    from portfolio_investimentos import migrations
    assert migrations.schema_version(app.config["DB_PATH"]) == migrations.LATEST
    assert not (Path(app.config["DB_PATH"]).parent / "copias").exists()  # nada a copiar numa base nova


def test_old_database_is_migrated_with_backup(tmp_path, monkeypatch):
    import sqlite3
    from portfolio_investimentos import migrations
    _legacy_db(tmp_path, monkeypatch)
    app = create_app({"TESTING": True})
    path = app.config["DB_PATH"]
    assert migrations.schema_version(path) == migrations.LATEST
    conn = sqlite3.connect(path)
    assert "result" in {r[1] for r in conn.execute('PRAGMA table_info("transaction")')}
    assert "last_quote" in {r[1] for r in conn.execute("PRAGMA table_info(instrument)")}
    assert conn.execute('SELECT count(*) FROM "transaction"').fetchone()[0] == 1  # os dados seguem lá
    conn.close()
    copies = list((tmp_path / "copias").glob("portfolio-*-migracao-v0-v3.db"))
    assert len(copies) == 1  # a cópia é da base antes de migrar (sem as colunas novas)
    old = sqlite3.connect(copies[0])
    assert "result" not in {r[1] for r in old.execute('PRAGMA table_info("transaction")')}
    old.close()


def test_failed_migration_rolls_back_and_names_the_backup(tmp_path, monkeypatch):
    from portfolio_investimentos import migrations
    _legacy_db(tmp_path, monkeypatch)

    def broken(conn):
        conn.execute('ALTER TABLE "transaction" ADD COLUMN extra TEXT')
        raise RuntimeError("falha de teste")
    monkeypatch.setattr(migrations, "MIGRATIONS", [(1, "passo avariado", broken)])
    monkeypatch.setattr(migrations, "LATEST", 1)
    with pytest.raises(migrations.MigrationError) as exc:
        create_app({"TESTING": True})
    assert "passo 1" in str(exc.value) and "copias" in str(exc.value)
    assert migrations.schema_version(tmp_path / "portfolio.db") == 0
    import sqlite3
    conn = sqlite3.connect(tmp_path / "portfolio.db")
    assert "extra" not in {r[1] for r in conn.execute('PRAGMA table_info("transaction")')}  # desfeito
    conn.close()


def test_database_from_newer_version_is_refused(tmp_path, monkeypatch):
    import sqlite3
    from portfolio_investimentos import migrations
    _legacy_db(tmp_path, monkeypatch)
    conn = sqlite3.connect(tmp_path / "portfolio.db")
    conn.execute(f"PRAGMA user_version = {migrations.LATEST + 1}")
    conn.commit()
    conn.close()
    monkeypatch.setenv("PORTFOLIO_DATA_DIR", str(tmp_path))
    with pytest.raises(migrations.MigrationError, match="mais recente"):
        create_app({"TESTING": True})


def test_automatic_backups_before_destructive_actions_and_restore(app, client):
    from portfolio_investimentos.services import backups
    path = app.config["DB_PATH"]
    token = csrf(client)
    post(client, "reset_data", token, confirm="APAGAR")  # sem dados: nada a proteger
    assert backups.list_backups(path) == []
    post(client, "add_transaction", token, date="2024-01-02", type="DEPOSIT", amount="1000", currency="EUR")
    post(client, "add_transaction", token, date="2024-01-03", type="BUY", instrument="ACME", quantity="5", price="10")
    csv = b"date,type,ticker,quantity,price,amount,fees,tax,currency,name\n2024-02-01,DEPOSIT,,,,50,,,EUR,\n"
    client.post("/api/upload_statement", headers={"X-CSRF-Token": token},
                data={"broker_key": "generic", "file": (io.BytesIO(csv), "a.csv")}, content_type="multipart/form-data")
    assert [b.label for b in backups.list_backups(path)] == ["Antes de importar"]
    post(client, "reset_data", token, confirm="APAGAR")
    labels = [b.label for b in backups.list_backups(path)]
    assert labels == ["Antes de apagar tudo", "Antes de importar"]
    with app.app_context():
        assert Transaction.query.count() == 0
    # restaurar a cópia feita antes de apagar tudo: os dados voltam, e fica uma cópia do estado vazio
    name = backups.list_backups(path)[0].name
    assert post(client, "restore_backup", token, name=name).status_code == 200
    with app.app_context():
        assert Transaction.query.count() == 3
    assert backups.list_backups(path)[0].label == "Antes de restaurar uma cópia"
    assert client.get("/definicoes").status_code == 200 and client.get("/carteira").status_code == 200
    # só nomes de cópias existentes: nada de caminhos
    assert post(client, "restore_backup", token, name="../portfolio.db").status_code == 400
    assert post(client, "restore_backup", token, name="portfolio-20200101-000000-x.db").status_code == 400


def test_backups_keep_only_the_latest_ten(tmp_path):
    import sqlite3
    from portfolio_investimentos.services import backups
    db_path = tmp_path / "portfolio.db"
    sqlite3.connect(db_path).execute("CREATE TABLE t (x)").connection.close()
    for _ in range(13):
        backups.make_backup(db_path, "manual")
    assert len(backups.list_backups(db_path)) == backups.KEEP
